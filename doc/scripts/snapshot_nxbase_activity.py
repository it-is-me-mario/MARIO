"""Snapshot reviewed native-to-NACE containment paths without editing nxbase.

Current review: EXIOBASE electricity technologies. Follow actual activity
parents and stop at the first official anchor; never expand a broad anchor.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook


def snapshot(workbook, pilot):
    sheet = load_workbook(workbook, read_only=True, data_only=False)['activity']
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    rows = [dict(zip(headers, row), row=i) for i, row in enumerate(iterator, 2)]
    by_extended = {r['classification']+' | '+r['name']: r for r in rows}
    by_key = {(r['classification'], r['short']): r for r in rows}
    decisions = {
        'powerBW': ('35.11', '35.12',
                    'Biomass and waste combines feedstocks. The current nxbase '
                    'renewables anchor does not establish that all waste is renewable.'),
        'powerNC': ('35.11', '35.12',
                    'The nxbase fossil anchor is explicitly an editorial hypothesis. '
                    'Unclassified generation is not counted as a confirmed subdivision.'),
    }
    items = []
    for native in pilot['native']:
        if native['node'][0] != 'EXIO3I' or not native['label'].startswith('Production of electricity'):
            continue
        _, namespace, short = native['nxbase_key'].split('/')
        current = by_key[namespace, short]
        if current['name'] != native['label']:
            raise ValueError('Native label changed: '+native['code'])
        nodes = [native['node']]
        names = [native['label']]
        steps = []
        seen = set()
        while current.get('parent'):
            if current['parent'] in seen:
                raise ValueError('Cyclic activity graph')
            seen.add(current['parent'])
            parent = by_extended[current['parent']]
            if parent['classification'] not in ['NXB', 'NACE']:
                raise ValueError('Unexpected classification in reviewed route')
            code = re.sub(r'^[A-Z]\.', '', parent['short']) if parent['classification'] == 'NACE' else parent['short']
            destination = ['NACE2.1' if parent['classification']=='NACE' else 'NXB', code]
            steps.append({'source':'S11', 'location':f"activity!A{current['row']}:H{current['row']}",
                          'note':current.get('notes') or 'Existing nxbase activity containment parent.',
                          'relation':'Left within right', 'source_relation':'Left within right',
                          'source_left':nodes[-1], 'source_right':destination,
                          'raw_type':'Governed nxbase containment parent'})
            nodes.append(destination); names.append(parent['name']); current=parent
            if parent['classification']=='NACE':
                break
        if nodes[-1][0]!='NACE2.1' or not re.fullmatch(r'\d{2}\.\d{2}', nodes[-1][1]):
            raise ValueError('Reviewed route must end at an explicit NACE leaf')
        decision=decisions.get(short)
        items.append({'native':native['node'], 'name':native['label'], 'nodes':nodes,
                      'names':names, 'steps':steps, 'reference':nodes[-1][1],
                      'status':'pending' if decision else 'reviewed',
                      'candidate_references':list(decision[:2]) if decision else [],
                      'review_note':decision[2] if decision else
                      'Technology containment checked against NACE Rev. 2.1 explanatory notes, classes 35.11 and 35.12.'})
    if len(items)!=12:
        raise ValueError('Expected all twelve native generation categories')
    return {'schema_version':1, 'as_of':'2026-10-10', 'items':items,
            'source':{'id':'S11','title':'nxbase · reviewed EXIOBASE electricity activity paths',
                      'publisher':'nxbase','filename':workbook.name,
                      'url':'nxbase-activity-refinements.json',
                      'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest(),
                      'retrieved':'2026-10-10',
                      'role':'Actual activity parent paths. Reviewed technology containment takes precedence over broad product/producer correspondences.'},
            'evidence':{'id':'S17','title':'Eurostat · NACE Rev. 2.1 explanatory notes',
                        'publisher':'Eurostat',
                        'url':'https://ec.europa.eu/eurostat/documents/3859598/21633320/KS-GQ-24-007-EN-N.pdf',
                        'retrieved':'2026-10-10','role':'Classes 35.11 and 35.12, printed pages 211–212. Non-renewable generation includes nuclear.'}}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,required=True)
    parser.add_argument('--pilot',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    result=snapshot(args.workbook,json.loads(args.pilot.read_text()))
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'reviewed':sum(i['status']=='reviewed' for i in result['items']),
                      'pending':sum(i['status']=='pending' for i in result['items'])}))
