"""Build installable offline concordances from the explorer's canonical compiler.

The committed compressed bundles are runtime assets, not database observations.
Run with no arguments after refreshing explorer classification sources. Optional
HS inputs snapshot the supplied EXIOBASE v3 / HS 2007 table, joined by exact
typed product aliases; missing aliases and empty HS cells stay unmapped.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
import csv
from datetime import date
import gzip
import hashlib
import json
from pathlib import Path
import sys

DOC = Path(__file__).resolve().parents[1]
ROOT = DOC.parent
sys.path.insert(0, str(ROOT))
from mario.database_profiles.concordances import CLASSIFICATIONS
from build_sectoral_scope import build


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def snapshot_hs(path, descriptions, inputs):
    from openpyxl import load_workbook

    products = json.loads((inputs / "commodity-correspondences.json").read_text())
    aliases = {}
    for row in products["records"]:
        if row["namespace"] == "EXIO3P":
            alias = "C_" + row["short"]
            if alias in aliases:
                raise ValueError("Ambiguous typed EXIOBASE product alias: " + alias)
            aliases[alias] = row["code"]
    with gzip.open(descriptions, "rt", encoding="utf-8") as stream:
        names = {r["code"]: r["desc"] for r in csv.DictReader(stream) if len(r["code"]) == 6}
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["Sheet1"]
    assert next(sheet.values) == ("EX_code", "HS2007")
    rows, excluded = [], set()
    for number, (alias, raw) in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if raw is None:
            continue
        code = str(raw).zfill(6)
        if code not in names:
            raise ValueError("Invalid HS 2007 code: " + code)
        if alias not in aliases:
            excluded.add(alias)
            continue
        rows.append(dict(native=aliases[alias], alias=alias, code=code, location=f"Sheet1!A{number}:B{number}"))
    book.close()
    result = dict(schema_version=1, as_of=date.today().isoformat(),
        rows=rows, node_names={"HS2007:"+code: names[code] for code in sorted({r['code'] for r in rows})},
        excluded_aliases=sorted(excluded),
        sources={
            "H01": dict(id="H01", title="EXIOBASE v3 products / HS 2007 correspondence, supplied with BONSAI",
                publisher="CML / BONSAI", filename=path.name, sha256=digest(path), url="https://bonsai.uno/",
                role="Original EXIOBASE v3 correspondence described in the supplied HS_vs_exio/readme.txt. Exact C_ product aliases are joined to typed EXIO3P identities, never to activity names. The homepage is attribution, not a verified download link. Binary links are qualified and supply no allocation weights."),
            "H02": dict(id="H02", title="HS 2007 code descriptions", publisher="UN Comtrade",
                filename=descriptions.name, sha256=digest(descriptions),
                url="https://comtrade.un.org/data/cache/classificationH3.json",
                role="UN Comtrade HS 2007 descriptions retained in the local concordance Python port. Used to validate six-digit codes and label results; not a product concordance.")})
    (inputs / "hs2007-exiobase.json").write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":"))+"\n")
    return result


def main(inputs, scope, output):
    source = build(inputs, json.loads(scope.read_text()))
    hs = json.loads((inputs / "hs2007-exiobase.json").read_text())
    source["sources"].update(hs["sources"])
    source["node_names"].update(hs["node_names"])
    official = json.loads((inputs / "official-activity-bridges.json").read_text())
    products = json.loads((inputs / "commodity-correspondences.json").read_text())
    legacy = json.loads((inputs / "reviewed-activity-correspondences.json").read_text())
    for data in (official, products, legacy):
        source["sources"].update({s["id"]: s for s in data["sources"]})
        source["node_names"].update(data.get("node_names", {}))
    source["node_names"].update({"NACE1.1:"+n["code"]: n["name"] for n in legacy["nace11_nodes"]})
    for ref in source["references"].values():
        namespace = "NACE2.1" if ref["short_label"] == "NACE" else "CPA2.2"
        source["node_names"].update({namespace+":"+n["code"]: n["name"] for n in ref["nodes"]})

    # Only official taxonomy-to-taxonomy bridges of the SAME semantic kind.
    # Producer edges never fabricate product membership from an industry.
    edges = [*official["edges"], *products["edges"], *legacy["legacy_revision_edges"],
             *({**step, "left": step["source_left"], "right": step["source_right"]}
               for step in source["route_steps"].values())]
    bridges = {}
    for edge in edges:
        left, right = edge["left"], edge["right"]
        if left[0] not in CLASSIFICATIONS or right[0] not in CLASSIFICATIONS:
            continue
        if CLASSIFICATIONS[left[0]][1] != CLASSIFICATIONS[right[0]][1]:
            continue
        if left[0] == right[0]:
            continue
        for reverse in (False, True):
            own = {**edge, "left": right if reverse else left, "right": left if reverse else right,
                   "source_left": left, "source_right": right, "direction": "reverse" if reverse else "forward",
                   "source_left_part": edge.get("left_part"), "source_right_part": edge.get("right_part"),
                   "left_part": edge.get("right_part" if reverse else "left_part"),
                   "right_part": edge.get("left_part" if reverse else "right_part")}
            key = (tuple(left), tuple(right), edge["source"], edge.get("location", ""), reverse)
            bridges.setdefault(key, own)
    bridges = list(bridges.values())
    grouped = defaultdict(list)
    for kind, columns in source["columns"].items():
        for column in columns:
            own = {k: column[k] for k in ("id", "family", "units", "namespace", "native", "releases", "formats", "source_url")}
            own.update(kind=kind, reference="NACE2.1" if kind == "activity" else "CPA2.2")
            grouped[column["family"]].append(own)
            if column["family"] == "BONSAI" and kind == "activity":
                grouped["BONSAI"].append({**own, "id": own["id"]+"-hiot-reference-process",
                    "native": source["country_variants"]["BONSAIHIOTA"], "reference_process": True})

    # Explicit product identities used by the explorer are the only HS joins.
    by_native = defaultdict(list)
    for row in hs["rows"]:
        step = dict(source_left=["EXIO3P",row["native"]], source_right=["HS2007",row["code"]],
                    source="H01", location=row["location"], relation="Supplied EXIOBASE product / HS 2007 correspondence",
                    note="Exact typed product alias "+row["alias"]+"; qualified binary footprint, not allocation weights.")
        by_native[row["native"]].append(step)
    for profile in grouped["EXIOBASE"]:
        if profile["kind"] != "commodity":
            continue
        for native in profile["native"]:
            # SUT/HIOT product reuse requires an explicit supplied bridge,
            # not merely a code that happens to share a spelling.
            prefixes = {native["code"]: []} if profile["namespace"] == "EXIO3P" else {}
            for link in native["links"].values():
                steps = [source["route_steps"][sid] for sid in link["steps"]]
                for i, step in enumerate(steps):
                    if step["source_right"][0] == "EXIO3P":
                        prefixes[step["source_right"][1]] = steps[:i+1]
            routes = []
            for code, prefix in prefixes.items():
                for step in by_native[code]:
                    steps = [*prefix, step]
                    routes.append(dict(nodes=[steps[0]["source_left"], *[s["source_right"] for s in steps]],
                        steps=steps, status="partial", mapping_origin="published_hs2007",
                        note="Supplied EXIOBASE v3 product / HS 2007 correspondence; six-digit codes validated against UN Comtrade. No allocation weights are supplied."))
            if routes:
                native["additional_routes"] = {"HS2007": routes}

    output.mkdir(parents=True, exist_ok=True)
    hierarchies = {ref["namespace"]: {"source": ref["source"], "nodes": ref["nodes"]}
                   for ref in source["references"].values()}
    shared = dict(bridges=bridges, node_names=source["node_names"], sources=source["sources"],
                  hierarchies=hierarchies)
    for own in shared["sources"].values():
        if own.get("url") and not own["url"].startswith(("https://", "http://")):
            own["url"] = "https://www.marioproject.com/en/latest/_static/data/economy/" + own["url"]
    (output / "shared.json.gz").write_bytes(gzip.compress(
        json.dumps(shared, ensure_ascii=False, separators=(",", ":")).encode(), compresslevel=9, mtime=0))
    manifest = dict(schema_version=1, as_of=source["as_of"],
        note="Compiled from the Database explorer's canonical classification sources. Qualified links are not allocation matrices; native parents are context only.",
        input_checksums={f.name: digest(f) for f in sorted(inputs.glob("*.json"))},
        families=sorted(grouped), profiles=[])
    for family, profiles in sorted(grouped.items()):
        used_steps, direct = set(), set()
        for profile in profiles:
            direct = set()
            for native in profile["native"]:
                direct.update(native.get("additional_routes", {}))
                for link in [*native["links"].values(), *native.get("coarse", {}).get("anchors", [])]:
                    for route in [link, *link.get("alternatives", []), *({"steps": s} for s in link.get("alternative_steps", []))]:
                        used_steps.update(route["steps"])
                        for sid in route["steps"]:
                            step = source["route_steps"][sid]
                            direct.update(n[0] for n in (step["source_left"],step["source_right"]) if n[0] in CLASSIFICATIONS)
            profile["direct_classifications"] = sorted(direct)
            reachable = {profile["reference"], *direct}
            for _ in CLASSIFICATIONS:
                expanded = reachable | {s["right"][0] for s in bridges if s["left"][0] in reachable}
                if expanded == reachable:
                    break
                reachable = expanded
            profile["classifications"] = [key for key in CLASSIFICATIONS if key in reachable]
            manifest["profiles"].append({k: v for k, v in profile.items() if k != "native"})
        steps = {sid: source["route_steps"][sid] for sid in sorted(used_steps)}
        bundle = dict(profiles=profiles, route_steps=steps)
        payload = json.dumps(bundle, ensure_ascii=False, separators=(",", ":")).encode()
        target = output / (family.lower()+".json.gz")
        target.write_bytes(gzip.compress(payload, compresslevel=9, mtime=0))
        print(f"{family}: {len(profiles)} profiles, {len(payload):,} uncompressed bytes, {target.stat().st_size:,} packaged bytes")
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n")


if __name__ == "__main__":
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("--inputs", type=Path, default=DOC / "classifications/sectoral")
    cli.add_argument("--scope", type=Path, default=DOC / "source/_static/data/economy/country-scope.json")
    cli.add_argument("--output", type=Path, default=ROOT / "mario/database_profiles/data/concordances")
    cli.add_argument("--hs", type=Path)
    cli.add_argument("--hs-descriptions", type=Path)
    args = cli.parse_args()
    if bool(args.hs) != bool(args.hs_descriptions):
        cli.error("Pass both --hs and --hs-descriptions to refresh the HS snapshot")
    if args.hs:
        snapshot_hs(args.hs, args.hs_descriptions, args.inputs)
    main(args.inputs, args.scope, args.output)
