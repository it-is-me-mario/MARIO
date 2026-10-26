from __future__ import annotations

from html import escape
import json
from pathlib import Path

from openpyxl import load_workbook


DOC_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = DOC_ROOT / "source" / "_static" / "data" / "supporting_files" / "Nomenclature_and_formulas.xlsx"
OUTPUT_DIR = DOC_ROOT / "source" / "get_started" / "_generated"


def normalize_rows(sheet_name: str):
    workbook = load_workbook(WORKBOOK, data_only=True)
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value).strip() for value in rows[0] if value not in (None, "")]
    width = len(headers)
    body = []
    for row in rows[1:]:
        values = list(row[:width])
        if all(value in (None, "") for value in values):
            continue
        body.append(["" if value is None else str(value).strip() for value in values])
    return headers, body


def render_table(
    table_id: str,
    headers: list[str],
    rows: list[list[str]],
    preview_rows: int | None = None,
    preview_label: str = "rows",
) -> str:
    """Render one filterable table.

    With ``preview_rows``, the page shows only the first rows and a toggle to
    expand the rest (see ``_static/terminology-tables.js``). Without JavaScript
    the whole table stays visible.
    """
    preview = ""
    if preview_rows:
        preview = f' data-preview-rows="{preview_rows}" data-preview-label="{escape(preview_label)}"'
    html = []
    html.append(f'<div class="terminology-table-wrapper">')
    html.append(f'  <table id="{table_id}" class="terminology-table"{preview}>')
    html.append("    <thead>")
    html.append("      <tr>")
    for header in headers:
        html.append(f"        <th>{escape(header)}</th>")
    html.append("      </tr>")
    html.append('      <tr class="terminology-column-filters">')
    for index, header in enumerate(headers):
        html.append("        <th>")
        html.append(
            f'          <input class="terminology-column-filter" data-col-index="{index}" '
            f'type="search" aria-label="Filter {escape(header)}" placeholder="Filter {escape(header)}..." />'
        )
        html.append("        </th>")
    html.append("      </tr>")
    html.append("    </thead>")
    html.append("    <tbody>")
    for row in rows:
        search_blob = " ".join(row).lower()
        html.append(f'      <tr data-search="{escape(search_blob)}">')
        for index, cell in enumerate(row):
            if headers and headers[0] == "Matrix" and index == 0:
                html.append(f"        <td><strong>{escape(cell)}</strong></td>")
            else:
                html.append(f"        <td>{escape(cell)}</td>")
        html.append("      </tr>")
    html.append("    </tbody>")
    html.append("  </table>")
    html.append("</div>")
    return "\n".join(html) + "\n"


# Layout and reading order are editorial; names, definitions and axes come from
# the same workbook as the complete reference. Keep the first view on native
# blocks; the full runtime catalogue remains linked from each format page.
FORMATS = {
    "iot": {
        "columns": ["Sectors", "Final demand"],
        "rows": ["Sectors", "Factors", "Satellite accounts"],
        "blocks": [["Z", "Y"], ["V", "VY"], ["E", "EY"]],
        "derived": ["X", "z", "v", "e", "w", "f", "F", "m", "M", "p"],
    },
    "sut": {
        "columns": ["Commodities", "Activities", "Final demand"],
        "rows": ["Commodities", "Activities", "Factors", "Satellite accounts"],
        "blocks": [[None, "U", "Yc"], ["S", None, "Ya"],
                   ["Vc", "Va", "VY"], ["Ec", "Ea", "EY"]],
        "derived": ["Xa", "Xc", "u", "s", "c", "va", "vc", "ea", "ec",
                    "waa", "wcc", "wac", "wca", "fa", "fc", "Fa", "Fc",
                    "ma", "mc", "pa", "pc"],
    },
}
BLOCK_LABELS = {"Z": "Transactions", "U": "Use", "S": "Supply", "Y": "Final demand",
                "V": "Factors", "E": "Satellite accounts"}



def nam_domain_label(schema, key):
    return " + ".join(schema["sets"][group]["label"]
                      for group in schema["domains"][key]["sets"]) or schema["domains"][key].get("label", "Balancing item")


def _mode_controls():
    return ('<div class="table-map__controls" data-table-mode-controls hidden '
            'role="group" aria-label="Matrix representation">'
            '<button type="button" data-table-mode-button="flows" aria-pressed="true">Flows</button>'
            '<button type="button" data-table-mode-button="coefficients" aria-pressed="false">Coefficients</button></div>')


def _map_panel(kind, mode, rows, columns, cells):
    """Generate complete static links/definitions for one representation."""
    coefficient = mode == "coefficients"
    caption = f'{kind.upper()} {"coefficients" if coefficient else "native flows"} · rows × columns'
    parts = [f'<div class="table-map__mode" data-table-mode="{mode}">',
             '<div class="table-map__scroll" tabindex="0" role="region" '
             f'aria-label="{kind.upper()} {mode} block diagram; scroll horizontally on small screens">',
             '<table class="table-map__grid">', f'<caption>{caption}</caption>', '<thead><tr><td></td>']
    for key, label, full in columns:
        attribute = f' data-set-id="{key}"' if kind == "nam" else ""
        text = f'<span class="table-map__column-label">{escape(label)} - {escape(full)}</span>' if kind == "nam" else escape(label)
        parts.append(f'<th scope="col"{attribute}>{text}</th>')
    parts.append('</tr></thead><tbody>')
    shown = []
    for r, (row_key, label, full) in enumerate(rows):
        attribute = f' data-set-id="{row_key}" title="{escape(full, quote=True)}"' if kind == "nam" else ""
        parts.append(f'<tr><th scope="row"{attribute}>{escape(label)}</th>')
        for c, (column_key, _, _) in enumerate(columns):
            block = cells[r][c]
            if block is None:
                parts.append('<td class="table-map__empty"><span aria-label="No native block in this representation">—</span></td>')
            elif block.get("unavailable"):
                parts.append(f'<td class="table-map__unavailable"><span title="No built-in coefficient; this final-demand block remains an exogenous flow">'
                             f'<code>{escape(block["flow"])}</code><small>Exogenous flow</small></span></td>')
            else:
                name = block["name"]
                shown.append(block)
                parts.append(f'<td class="table-map__block is-{block["family"]}">'
                             f'<a href="#matrix-{name}" data-matrix-link="matrix-{name}" '
                             f'data-flow-name="{block["flow"]}" data-row-domain="{row_key}" data-column-domain="{column_key}" '
                             f'aria-label="{name}: {escape(block["description"], quote=True)}">'
                             f'<code>{name}</code>{block.get("sublabel", "")}</a></td>')
        parts.append('</tr>')
    parts.append('</tbody></table></div><div class="table-map__definitions" aria-live="polite">')
    for block in shown:
        name = block["name"]
        parts.append(f'<section id="matrix-{name}" class="table-map__definition" data-matrix-definition>'
                     f'<p class="table-map__symbol"><code>{name}</code><span>{"Coefficients" if coefficient else "Native flows"}</span></p>'
                     f'<p class="table-map__description">{escape(block["description"])}</p><dl>'
                     f'<dt>Rows</dt><dd>{escape(block["rows"])}</dd>'
                     f'<dt>Columns</dt><dd>{escape(block["columns"])}</dd>')
        if block.get("denominator"):
            parts.append(f'<dt>Column denominator</dt><dd><code>{block["denominator"]}</code> · full account expenditure</dd>')
        parts.append('</dl></section>')
    parts.append('</div></div>')
    return '\n'.join(parts)


def render_nam_structure(schema):
    """Show single-set flow/coefficient pairs under one import authority."""
    axes = [(key, value["short_name"], value["label"]) for key, value in schema["sets"].items()]
    alternative = schema["representations"]["use_imports"]["alternative_import_blocks"]
    blocks = {(m["rows"], m["columns"]): m for m in schema["matrices"]
              if m["role"] == "flow" and m["name"] not in alternative}
    coefficients = {m["source"]: m for m in schema["matrices"] if m["role"] == "coefficient"}
    panels = []
    for mode in ("flows", "coefficients"):
        cells = []
        for row_key, _, _ in axes:
            row = []
            for column_key, _, _ in axes:
                flow = blocks.get((row_key, column_key))
                if flow is None:
                    row.append(None)
                    continue
                block = coefficients[flow["name"]] if mode == "coefficients" else flow
                family = ("trade" if row_key == "external_trade" else
                          "demand" if column_key in ("consumption", "nonfinancial_asset") else
                          "capital" if any("capital" in key for key in (row_key, column_key)) else
                          "factors" if row_key in ("factor", "tax", "income_resident", "income_nonresident", "distribution") else "core")
                row.append({**block, "flow": flow["name"], "family": family,
                            "rows": nam_domain_label(schema, row_key), "columns": nam_domain_label(schema, column_key)})
            cells.append(row)
        panels.append(_map_panel("nam", mode, axes, axes, cells))
    return ('<div class="table-map table-map--nam" data-table-map="nam">' + _mode_controls()
            + '\n'.join(panels) + '</div>\n')


def render_structure(kind, records):
    layout = FORMATS[kind]
    rows = [(label, label, label) for label in layout["rows"]]
    columns = [(label, label, label) for label in layout["columns"]]
    panels = []
    for mode in ("flows", "coefficients"):
        cells = []
        for native in layout["blocks"]:
            row = []
            for flow in native:
                name = flow.lower() if flow and mode == "coefficients" else flow
                if flow is None:
                    row.append(None)
                elif name not in records:
                    row.append({"unavailable": True, "flow": flow})
                else:
                    record = records[name]
                    family = {"Y": "demand", "V": "factors", "E": "satellite"}.get(flow[0], "core")
                    row.append({"name": name, "flow": flow, "family": family,
                                "description": record["Description"].rstrip("*"), "rows": record["Row indices"],
                                "columns": record["Column indices"], "sublabel": f'<span>{BLOCK_LABELS[flow[0]]}</span>'})
            cells.append(row)
        panels.append(_map_panel(kind, mode, rows, columns, cells))
    return (f'<div class="table-map" data-table-map="{kind}">' + _mode_controls()
            + '\n'.join(panels) + '</div>\n')


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    nam = json.loads((DOC_ROOT.parent / "mario/model/nam_schema.json").read_text())
    (OUTPUT_DIR / "nam_structure.html").write_text(render_nam_structure(nam), encoding="utf-8")
    (OUTPUT_DIR / "nam_indices.html").write_text(render_table(
        "nam-indices", ["Short name", "Set ID", "Display label", "Account total", "Classification"],
        [[value["short_name"], key, value["label"], value["total_vector"], value["classification"]] for key, value in nam["sets"].items()],
        preview_rows=6, preview_label="sets"), encoding="utf-8")
    (OUTPUT_DIR / "nam_matrices.html").write_text(render_table(
        "nam-matrices", ["Matrix", "Rows", "Columns", "Role", "Description"],
        [[m["name"], nam_domain_label(nam, m["rows"]), nam_domain_label(nam, m["columns"]), m["role"], m["description"]]
         for m in nam["matrices"]], preview_rows=8, preview_label="matrices"), encoding="utf-8")

    indices_headers, indices_rows = normalize_rows("Indices")
    matrices_headers, matrices_rows = normalize_rows("Matrices")
    records = {row[0]: dict(zip(matrices_headers, row)) for row in matrices_rows}

    for kind, layout in FORMATS.items():
        (OUTPUT_DIR / f"{kind}_structure.html").write_text(render_structure(kind, records), encoding="utf-8")
        indices = [[row[0], row[3]] for row in indices_rows if row[1] in ("All", kind.upper())]
        (OUTPUT_DIR / f"{kind}_indices.html").write_text(
            render_table(f"{kind}-indices", ["Index", "Description"], indices), encoding="utf-8")
        columns = ["Matrix", "Description", "Row indices", "Column indices"]
        derived = [[records[name][column] for column in columns] for name in layout["derived"]]
        (OUTPUT_DIR / f"{kind}_derived.html").write_text(
            render_table(f"{kind}-derived", columns, derived, preview_rows=5, preview_label="matrices"),
            encoding="utf-8")

    (OUTPUT_DIR / "indices_table.html").write_text(
        render_table(
            table_id="terminology-indices-table",
            headers=indices_headers,
            rows=indices_rows,
        ),
        encoding="utf-8",
    )

    (OUTPUT_DIR / "matrices_table.html").write_text(
        render_table(
            table_id="terminology-matrices-table",
            headers=matrices_headers,
            rows=matrices_rows,
            preview_rows=4,
            preview_label="matrices",
        ),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
