"""Inventory source-published scope for databases without a MARIO parser.

Extract only labels, identifiers and source metadata; never emission-factor
values or economic matrices. Keep these profiles outside the parser catalogue.
"""
from pathlib import Path
import argparse, hashlib, json
import openpyxl


RELEASE_NOTES={'2025':'https://cdn.sanity.io/files/3ogo9b9g/production/2bdd4a975ba09b579c9859c6d2879702e19eccd9.pdf'}


def read_open_ceda(book):
    """Release, base year, licence, release date, sectors, regions and source ranges of either workbook layout."""
    if 'Open CEDA' in book.sheetnames:  # 2024 layout: one sheet with the factors and their metadata
        sheet=book['Open CEDA']
        headers=list(sheet.iter_rows(min_row=27,max_row=28,values_only=True))
        return {'release':str(sheet['D13'].value).removeprefix('CEDA ').strip(),'year':int(sheet['C24'].value),
                'licence':str(sheet['D14'].value).strip(),'release_date':str(sheet['D12'].value),
                'sectors':[{'code':str(code).strip(),'label':str(label).strip()} for code,label in zip(headers[1][4:],headers[0][4:]) if code is not None],
                'regions':[{'code':str(r[1]).strip(),'label':str(r[2]).strip()} for r in sheet.iter_rows(min_row=29,max_col=3,values_only=True) if r[1]],
                'ranges':{'Open CEDA':['B12:D14','B24:C26','E27:ON28','B29:C177'],'Metadata':['A1:B401']}}
    # 2025 layout: release metadata on the cover; base-year factors (USD) in GHG_t_Raw.
    cover={str(r[1]).strip():r[3] for r in book['Cover'].iter_rows(max_col=4,values_only=True) if r[1] in ('Release date','Version','License')}
    rows=list(book['GHG_t_Raw'].iter_rows(values_only=True))
    if rows[0][0]!='Year' or rows[3][0]!='Country Code':
        raise ValueError('Unexpected GHG_t_Raw layout')
    sectors=[{'code':str(code).strip(),'label':str(label).strip()} for code,label in zip(rows[3][3:],rows[2][3:]) if code is not None]
    regions=[{'code':str(r[0]).strip(),'label':str(r[1]).strip()} for r in rows[4:] if r[0]]
    last=openpyxl.utils.get_column_letter(3+len(sectors))
    return {'release':str(cover['Version']).removeprefix('CEDA ').strip(),'year':int(rows[0][1]),'licence':str(cover['License']).strip(),
            'release_date':str(cover['Release date']).split(' ')[0],'sectors':sectors,'regions':regions,
            'ranges':{'Cover':['B9:D11'],'GHG_t_Raw':['A1:B1',f'D3:{last}4',f'A5:B{4+len(regions)}'],'Metadata':['A1:B401']}}


def build_open_ceda(path, data):
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        source=read_open_ceda(book)
        release,year,sectors,regional=source['release'],source['year'],source['sectors'],source['regions']
        if source['licence']!='CC BY-SA 4.0':
            raise ValueError('Check the source licence before updating this profile')
        metadata={str(row[0]).strip():str(row[1]).strip() for row in book['Metadata'].iter_rows(min_row=2,max_col=2,values_only=True) if row[0]}
        if len(sectors)!=400 or {s['code']:s['label'] for s in sectors}!=metadata:
            raise ValueError('Sector identifiers/names do not match the source metadata')
        if len(regional)!=149 or len({r['code'] for r in regional})!=149:
            raise ValueError('Unexpected or duplicate CEDA regions')
        world={c['iso3']:c for c in json.loads((data/'world-country-map.json').read_text())['countries']}
        regions=[]
        for r in regional:
            if r['code']=='ROW':
                regions.append({'label':'ROW','name':r['label'],'kind':'aggregate','members':[],'membership_status':'not_mapped'})
            elif r['code'] in world:
                regions.append({'label':r['code'],'name':world[r['code']]['name'],'source_name':r['label'],'kind':'country','iso3':r['code']})
            else:
                raise ValueError(f"Unreviewed CEDA regional code: {r['code']}")
        profile_id=f'open-ceda-{release}-{year}'
        sets={'Region':[r['code'] for r in regional],'Sector':[s['label'] for s in sectors]}
        digest=lambda values:hashlib.sha256(json.dumps(values,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
        profile={'schema_version':1,'id':profile_id,'source':'CEDA','guide':'ceda','label':f'Open CEDA {release} by Watershed · base year {year}',
                 'release':release,'year':year,'years':[year],'table':'Emission factors','status':'source_verified','parser_supported':False,
                 'environmental':True,'counts':{'Region':len(regions),'Sector':len(sectors)},'sets':sets,'native_codes':{s['label']:s['code'] for s in sectors},
                 'region_counts':[{'region':r['label'],'set':'Sector','count':len(sectors)} for r in regions],
                 'region_items':[],'set_flags':{'Satellite account':{'status':'not_provided'},'Factor of production':{'status':'not_provided'},'Consumption category':{'status':'not_provided'}},
                 'geography':{'status':'source_verified_individual_regions','scope':'global','regions':regions,'individual_countries':sum(r['kind']=='country' for r in regions),
                              'aggregate_count':sum(r['kind']=='aggregate' for r in regions),'unresolved_count':0,'country_memberships':[],
                              'method':'Native ISO3 identifiers verified against the country registry; ROW retained as a global aggregate.'},
                 'coverage_signatures':{'geography':digest(sorted(sets['Region'])),'economic':digest(sectors)},
                 'classifications':{'status':'native_labels','native_classification':'CEDA sectors based on BEA. NAICS/NACE correspondences remain unverified.','mappings':[]},
                 'provenance':{'method':'source_metadata','source_filename':path.name,'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                               'source_url':'https://watershed.com/en-GB/solutions/ceda','methodology_url':'https://cdn.sanity.io/files/3ogo9b9g/production/70c90f1ce6449776c2987aafe205bf6239952776.pdf',
                               'source_release_date':source['release_date'],'source_ranges':source['ranges'],**({'release_note_url':RELEASE_NOTES[release]} if release in RELEASE_NOTES else {}),
                               'options':{'unit':'Monetary'},'reading_scope':'Source metadata only; no MARIO parse or indicator values.'},
                 'source_url':'https://watershed.com/en-GB/solutions/ceda','license_url':'https://creativecommons.org/licenses/by-sa/4.0/',
                 'attribution':'CEDA by Watershed. Metadata adapted from Open CEDA under CC BY-SA 4.0.',
                 'scope_note':'Open CEDA publishes characterized GHG emission factors'+(' (IPCC AR6 GWP100 from CEDA 2025)' if int(release)>=2025 else '')+'. Raw satellite accounts, factors of production and final-demand categories are not supplied in this workbook. MARIO has no CEDA parser; indicators remain unavailable.',
                 'workbook':f'_static/data/external_profiles/{profile_id}-sets.xlsx','detail':f'_static/data/external_profiles/{profile_id}.json'}
        return profile
    finally:
        book.close()


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--open-ceda',type=Path,required=True);p.add_argument('--data',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    args=p.parse_args();profile=build_open_ceda(args.open_ceda,args.data);args.output.mkdir(parents=True,exist_ok=True)
    (args.output/(profile['id']+'.json')).write_text(json.dumps(profile,ensure_ascii=False,indent=2)+'\n')
    compact={k:v for k,v in profile.items() if k not in ['sets','native_codes','region_items']}
    registry=args.data/'external_profiles.json'
    previous=json.loads(registry.read_text())['profiles'] if registry.is_file() else []
    profiles=[p for p in previous if p['id']!=compact['id']]+[compact]
    (args.output/'external_profiles.json').write_text(json.dumps({'schema_version':1,'profiles':profiles},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'id':profile['id'],'counts':profile['counts'],'countries':profile['geography']['individual_countries'],'aggregates':profile['geography']['aggregate_count'],'parser_supported':False}))
