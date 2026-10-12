"""Snapshot EXIOBASE Hybrid activity parents from nxbase without editing its sets.

Leaf containment, broad containers and unresolved anchors stay distinct. Neither
an official parent nor a shared monetary label enumerates a native footprint.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook


def normalize_classifications(path):
    """Retain only the publisher's activity identity table, not economic data."""
    raw=path.read_bytes()
    if hashlib.md5(raw).hexdigest()!='5b50b2835b6a31c3fea8a959e3b697fe':
        raise ValueError('Review the source classifications release before replacing this snapshot')
    book=load_workbook(path,read_only=True,data_only=True)
    activities={}
    for i,row in enumerate(book['Activities'].values,1):
        if i==1:continue
        _,name,code,alias=row
        if name in activities:
            if (code,alias)!=(activities[name]['code'],activities[name]['alias']):
                raise ValueError('Country-specific activity identities differ')
        else:activities[name]={'name':name,'code':code,'alias':alias,'location':f'Activities!B{i}:D{i}'}
    book.close()
    if len(activities)!=164:raise ValueError('Expected 164 Hybrid activities')
    return {'schema_version':1,'source':{'id':'S23','publisher':'EXIOBASE','title':'EXIOBASE Hybrid 3.3.18 activity classifications',
        'url':'https://zenodo.org/records/7244919/files/Classifications_v_3_3_18.xlsx?download=1',
        'filename':'Classifications_v_3_3_18.xlsx','sha256':hashlib.sha256(raw).hexdigest(),
        'retrieved':'2026-10-10','role':'Published native activity codes and labels; no activity/product identity implied.'},
        'activities':list(activities.values())}


def snapshot(workbook,identities,reference,scope):
    book=load_workbook(workbook,read_only=True,data_only=False)
    values=list(book['activity'].values);book.close();headers=values[0]
    rows=[dict(zip(headers,r,strict=True),row=i) for i,r in enumerate(values,1) if i>3 and r[0] and r[3]]
    by_name={r['classification']+' | '+r['name']:r for r in rows}
    if len(by_name)!=len(rows):raise ValueError('Duplicate activity identity')
    natives={r['name']:r for r in rows if r['classification']=='EX3h'}
    official={a['name']:a for a in identities['activities']}
    profiles=[p for p in scope['profiles'] if p['family']=='EXIOBASE' and p['units']=='Hybrid' and 'Activity' in p['sets']]
    if not profiles or set(natives)!=set(official) or any(set(p['sets']['Activity'])!=set(natives) for p in profiles):
        raise ValueError('Hybrid parsed/publisher/nxbase activity vocabulary mismatch')
    ref={n['code']:n for n in reference['nodes']};leaves={c for c in ref if re.fullmatch(r'\d{2}\.\d{2}',c)}
    decisions={
        'powerBW':'Biomass and waste combines feedstocks; its renewable parent does not establish that all waste is renewable.',
        'powerNC':'The fossil parent is an editorial hypothesis; unclassified generation is not a confirmed subdivision.'}
    names={};items=[]
    for name,a in official.items():
        current=natives[name];nodes=[['EXIO3H',a['code']],['EX3h',current['short']]]
        names['EXIO3H:'+a['code']]=name;names['EX3h:'+current['short']]=name
        steps=[{'source':'S23','location':a['location'],'relation':'Native identity',
            'source_relation':'Native identity','source_left':nodes[0],'source_right':nodes[1],'note':f"nxbase activity!A{current['row']}:H{current['row']}; short {current['short']}; exact activity code/label identity verified"}]
        seen=set()
        while current.get('parent'):
            extended=current['parent']
            if extended in seen or extended not in by_name:raise ValueError('Cycle or missing activity parent: '+extended)
            seen.add(extended);parent=by_name[extended]
            if parent['classification'] not in ['NXB','NACE','NACE11']:
                raise ValueError('Unexpected cross-classification activity route')
            namespace={'NACE':'NACE2.1','NACE11':'NACE1.1'}.get(parent['classification'],parent['classification'])
            code=re.sub(r'^[A-Z]\.', '',parent['short']) if namespace.startswith('NACE') else parent['short']
            destination=[namespace,code];names[':'.join(destination)]=ref[code]['name'] if namespace=='NACE2.1' and code in ref else parent['name']
            steps.append({'source':'S24','location':f"activity!A{current['row']}:H{current['row']}",
                'relation':'Left within right','source_relation':'Left within right','source_left':nodes[-1],
                'source_right':destination,'note':current.get('notes') or 'Existing nxbase activity containment parent.'})
            nodes.append(destination);current=parent
            if namespace=='NACE2.1':break
        route={'nodes':nodes,'steps':steps,'relation':'Native activity containment declared in nxbase',
            'mapping_origin':'governed_nxbase','review_note':'Existing activity parent chain. A broad container is not an enumeration of native member classes.'}
        item={'code':a['code'],'name':name,'nxbase_short':natives[name]['short'],'links':{}}
        anchor=nodes[-1]
        if natives[name]['short'] in decisions:
            item['pending']={**route,'status':'pending','references':['35.11','35.12'],'review_note':decisions[natives[name]['short']]}
        elif anchor[0]=='NACE2.1' and anchor[1] in leaves:
            item['links'][anchor[1]]={**route,'status':'reviewed' if 'Production of electricity' in name else 'governed'}
        elif anchor[0]=='NACE2.1':
            code=anchor[1]
            if code not in ref:raise ValueError('Unknown NACE anchor '+code)
            candidates=sorted(c for c in leaves if c.startswith(code) and (code.isalpha() or c[:2]==code[:2]))
            if code.isalpha():
                # Official headings are ordered by section in the snapshot.
                section=None;sections={}
                for c in ref:
                    if c.isalpha():section=c
                    sections[c]=section
                candidates=sorted(c for c in leaves if sections[c]==code)
            sections=set();section=None
            for c in ref:
                if c.isalpha():section=c
                if c in candidates:sections.add(section)
            item['coarse']={'classification':'NACE Rev. 2.1','direct_nace':True,'reference_candidates':candidates,
                'reference_sections':sorted(sections),'note':'Actual nxbase broad containment anchor; leaf membership is unverified and excluded from the colour index.',
                'anchors':[{**route,'status':'coarse','code':code,'name':ref[code]['name'],'reference_context':{s:{} for s in sections}}]}
        else:
            item['unresolved']={**route,'status':'unresolved','review_note':'No current NACE Rev. 2.1 anchor in this nxbase activity chain.'}
        items.append(item)
    return {'schema_version':1,'as_of':'2026-10-10','namespace':'EXIO3H','classification':'EX3h','items':items,'node_names':names,
        'sources':[identities['source'],{'id':'S24','publisher':'nxbase','title':'nxbase · EXIOBASE Hybrid activity parent chains',
            'url':'nxbase-exiobase-hybrid.json','filename':workbook.name,'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest(),
            'retrieved':'2026-10-10','role':'Read-only snapshot of the existing EX3h activity parent chains. No parent expansion or commodity traversal.'}]}


if __name__=='__main__':
    cli=argparse.ArgumentParser(description=__doc__)
    cli.add_argument('--workbook',type=Path,required=True)
    cli.add_argument('--classifications',type=Path,help='Optional publisher workbook for the initial identity snapshot')
    cli.add_argument('--doc',type=Path,default=Path(__file__).resolve().parents[1]);args=cli.parse_args()
    inputs=args.doc/'classifications/sectoral';identities_path=inputs/'exiobase-hybrid-activities.json'
    if args.classifications:
        identities=normalize_classifications(args.classifications);identities_path.write_text(json.dumps(identities,ensure_ascii=False,indent=2)+'\n')
    else:identities=json.loads(identities_path.read_text())
    result=snapshot(args.workbook,identities,json.loads((inputs/'nace21-structure.json').read_text()),
        json.loads((args.doc/'source/_static/data/economy/country-scope.json').read_text()))
    for target in [inputs/'nxbase-exiobase-hybrid.json',args.doc/'source/_static/data/economy/nxbase-exiobase-hybrid.json']:
        target.write_text(json.dumps(result,ensure_ascii=False,separators=(',',':'))+'\n')
    print(json.dumps({'activities':len(result['items']),'leaf_anchors':sum(bool(i['links']) for i in result['items']),
        'broad_anchors':sum('coarse' in i for i in result['items']),'pending':sum('pending' in i for i in result['items']),
        'unresolved':sum('unresolved' in i for i in result['items'])}))
