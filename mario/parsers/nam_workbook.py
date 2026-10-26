"""Versioned native NAM workbooks, shared by templates, export and parsing."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mario.internal import ModelState, ModelStateMetadata
from mario.internal.nam import NAMSchema
from mario.log_exc.exceptions import WrongInput
from mario.model.nam import NAM_BLOCKS, NAM_SETS, nam_default_layout, nam_set_id


FORMAT = "MARIO native NAM"
VERSION = 2
ACCOUNT_COLUMNS = ("Region", "Set", "Item", "kind", "description", "metadata")
REGISTRY = ("scenario", "matrix", "presence", "sheet", "start_row", "start_column",
            "rows", "columns", "row_names", "column_names", "metadata")


def _write_row(sheet, row, values, column=1):
    for offset, value in enumerate(values):
        cell = sheet.cell(row, column + offset, value)
        if isinstance(value, str):
            cell.data_type = "s"


def _style_header(sheet, row, columns):
    for cell in sheet.iter_rows(min_row=row, max_row=row, max_col=columns):
        for item in cell:
            item.fill = PatternFill("solid", fgColor="245D65")
            item.font = Font(color="FFFFFF", bold=True)
            item.alignment = Alignment(vertical="top", wrap_text=True)


def write_nam_workbook(state, path):
    """Write local native records; derived accounting views/results stay outputs."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    flows = wb.active
    flows.title = "flows"
    _write_row(flows, 1, ["NAM native flows. Rows receive payments; columns pay."])
    registry = wb.create_sheet("matrix_layouts")
    _write_row(registry, 1, REGISTRY)
    accounts = wb.create_sheet("accounts")
    _write_row(accounts, 1, ACCOUNT_COLUMNS)
    for row, node in enumerate(state.nam_schema.accounts(), 2):
        region, label, item = node
        attributes = dict(state.nam_schema.item_metadata.get(region, {}).get(nam_set_id(label), {}).get(item, {}))
        kind, description = attributes.pop("kind", None), attributes.pop("description", None)
        _write_row(accounts, row, [*node, kind, description, json.dumps(attributes, ensure_ascii=False)])
    descriptors = wb.create_sheet("descriptors")
    _write_row(descriptors, 1, ["role", "item"])
    row = 2
    for role, items in state.nam_schema.descriptors.items():
        for item in items:
            _write_row(descriptors, row, [role, item])
            row += 1
    meta = wb.create_sheet("metadata")
    _write_row(meta, 1, ["key", "value"])
    configuration = state.nam_schema.to_dict()
    metadata = {
        "format": FORMAT, "version": VERSION,
        **{key: configuration[key] for key in ("representation", "currency", "scale", "period")},
        "name": state.metadata.name, "source": state.metadata.source,
        "year": state.metadata.year, "price": state.metadata.price,
        "license": state.metadata.license, "database_version": state.metadata.version,
        "history": json.dumps(state.metadata.history, ensure_ascii=False),
        "layouts": json.dumps(configuration["matrix_layouts"]),
        "scenarios": json.dumps({name: item.parent for name, item in state.scenarios.items()}),
    }
    extra = json.dumps(state.metadata.extra, ensure_ascii=False, allow_nan=False)
    for offset in range(0, len(extra), 30000):
        metadata[f"extra.{offset // 30000:04}"] = extra[offset:offset + 30000]
    for row, pair in enumerate(metadata.items(), 2):
        _write_row(meta, row, pair)

    cursor = 4
    for scenario, local in state.scenarios.items():
        for name, record in local.blocks.items():
            if name not in NAM_BLOCKS or NAM_BLOCKS[name]["role"] != "flow":
                continue
            presence = record.metadata["presence"]
            value = None if presence == "unavailable" else state.get_block(name, scenario)
            if value is None:
                if state.nam_schema.layout(name) == nam_default_layout(name):
                    skeleton = state.nam_schema.zero_block(name)
                    index, columns = skeleton.index, skeleton.columns
                else:
                    layout = state.nam_schema.layout(name)
                    axes = [tuple("Region" if r == "reference_region" else r for r in layout[side][:-1]) + ("Set", "Item")
                            for side in ("rows", "columns")]
                    index = pd.MultiIndex.from_tuples([], names=axes[0])
                    columns = pd.MultiIndex.from_tuples([], names=axes[1])
            else:
                index, columns = value.index, value.columns
            row_levels, col_levels = index.nlevels, columns.nlevels
            if cursor + col_levels + len(index) > 1048576 or row_levels + len(columns) > 16384:
                raise WrongInput("NAM block exceeds Excel worksheet limits; reduce the selected reference economies.")
            _write_row(flows, cursor - 1, [f"{scenario}: {name} ({presence})"])
            flows.cell(cursor - 1, 1).font = Font(bold=True, color="245D65")
            for level in range(col_levels):
                _write_row(flows, cursor + level, columns.get_level_values(level).tolist(), row_levels + 1)
                _style_header(flows, cursor + level, row_levels + len(columns))
            _write_row(flows, cursor + col_levels - 1, list(index.names))
            for offset, node in enumerate(index):
                _write_row(flows, cursor + col_levels + offset, node)
            if value is not None:
                data = value.to_numpy(dtype=float)
                for offset, values in enumerate(data):
                    _write_row(flows, cursor + col_levels + offset, values.tolist(), row_levels + 1)
            reg_row = registry.max_row + 1
            _write_row(registry, reg_row, [scenario, name, presence, "flows", cursor, 1,
                                        len(index), len(columns), json.dumps(list(index.names)),
                                        json.dumps(list(columns.names)), json.dumps(record.metadata, ensure_ascii=False)])
            registry.cell(reg_row, 2).hyperlink = f"#'flows'!A{cursor - 1}"
            cursor += col_levels + len(index) + 3
    instructions = wb.create_sheet("instructions")
    for row, line in enumerate([
        "Fill a block completely, including zeros, then set its presence to observed in matrix_layouts.",
        "Use declared_zero only when the whole block is known to be zero. Unavailable means unknown, not zero.",
        "Change values on flows; account identities and layout ranges must agree with accounts and matrix_layouts.",
        "Parse with mario.parse_from_excel(path, table='NAM', mode='flows').",
        "T, account totals, lowercase coefficients and B9_d/B9_f are computed outputs, not additional workbook payments.",
        "Every native block has exactly one accounting set on each axis, with its complete membership.",
        "Use accounts kind/description for item types. Other accounts hold residual routing/adjustments; imports have External trade account.",
    ], 1):
        _write_row(instructions, row, [line])
    for sheet in wb:
        if sheet not in (flows, instructions):
            _style_header(sheet, 1, sheet.max_column)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for column in range(1, min(sheet.max_column, 8) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 24 if column != 2 else 46
    flows.freeze_panes = "D7"
    instructions.column_dimensions["A"].width = 115
    meta.column_dimensions["B"].width = 85
    registry.column_dimensions["K"].width = 55
    wb.save(output)
    return str(output)


def _key_values(sheet):
    result = {}
    for key, value, *_ in sheet.iter_rows(min_row=2, values_only=True):
        if key is None:
            continue
        if key in result:
            raise WrongInput(f"Duplicate NAM metadata key: {key}.")
        result[key] = value
    return result


def read_nam_workbook(path):
    """Read the versioned native format, validating every observed cell."""
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        required = {"metadata", "accounts", "descriptors", "matrix_layouts", "flows"}
        if not required <= set(wb.sheetnames):
            raise WrongInput(f"Not a MARIO NAM workbook: missing sheets {sorted(required - set(wb.sheetnames))}.")
        meta = _key_values(wb["metadata"])
        if meta.get("format") != FORMAT or meta.get("version") != VERSION:
            raise WrongInput("Unsupported MARIO NAM workbook format/version. The single-set schema requires version 2; reparse the original FIGARO source or create a new NAM template. Version 1 workbooks are not silently reclassified.")
        if tuple(next(wb["accounts"].iter_rows(values_only=True))[:len(ACCOUNT_COLUMNS)]) != ACCOUNT_COLUMNS:
            raise WrongInput("Invalid NAM accounts header: version 2 requires item metadata columns.")
        members, item_metadata = {}, {}
        for region, group, item, kind, description, attributes, *_ in wb["accounts"].iter_rows(min_row=2, values_only=True):
            if region is None and group is None and item is None:
                continue
            members.setdefault(region, {}).setdefault(group, []).append(item)
            attributes = json.loads(attributes) if attributes else {}
            if kind is not None:
                attributes["kind"] = kind
            if description is not None:
                attributes["description"] = description
            if attributes:
                item_metadata.setdefault(region, {}).setdefault(group, {})[item] = attributes
        descriptors = {}
        for role, item, *_ in wb["descriptors"].iter_rows(min_row=2, values_only=True):
            if role is not None:
                descriptors.setdefault(role, []).append(item)
        schema = NAMSchema(members, representation=meta["representation"], currency=meta["currency"],
                           scale=meta["scale"], period=meta["period"], descriptors=descriptors,
                           matrix_layouts=json.loads(meta["layouts"]), item_metadata=item_metadata)
        metadata = ModelStateMetadata("NAM", **{key: meta.get(key) for key in ("name", "source", "year", "price", "license")},
                                      version=meta.get("database_version"), history=json.loads(meta["history"]),
                                      extra=json.loads("".join(meta[key] for key in sorted(meta) if key.startswith("extra."))))
        state = ModelState(metadata, nam_schema=schema)
        parents = json.loads(meta["scenarios"])
        if parents.get("baseline", "missing") is not None:
            raise WrongInput("NAM workbook requires a baseline scenario without a parent.")
        pending = {name: parent for name, parent in parents.items() if name != "baseline"}
        while pending:
            ready = [name for name, parent in pending.items() if parent is None or parent in state.scenarios]
            if not ready:
                raise WrongInput("Unknown parent or cycle in NAM workbook scenarios.")
            for name in ready:
                state.create_scenario(name, pending.pop(name))
        rows = list(wb["matrix_layouts"].iter_rows(values_only=True))
        if tuple(rows[0][:len(REGISTRY)]) != REGISTRY:
            raise WrongInput("Invalid NAM matrix_layouts header.")
        seen, rectangles = set(), []
        # Read each sheet once: read-only openpyxl random access is quadratic.
        sheets = {name: list(wb[name].iter_rows(values_only=True)) for name in {r[3] for r in rows[1:] if r[0] is not None}}
        for raw in rows[1:]:
            if raw[0] is None:
                continue
            scenario, name, status, sheet, top, left, nrows, ncols, row_names, col_names, details = raw[:len(REGISTRY)]
            if (scenario, name) in seen or scenario not in state.scenarios:
                raise WrongInput(f"Duplicate block or unknown scenario: {scenario}/{name}.")
            seen.add((scenario, name))
            if name not in NAM_BLOCKS or NAM_BLOCKS[name]["role"] != "flow":
                raise WrongInput(f"Only declared native NAM blocks are workbook inputs: {name}.")
            if any(isinstance(v, bool) or not isinstance(v, int) for v in (top, left, nrows, ncols)) or min(top, left) < 1 or min(nrows, ncols) < 0:
                raise WrongInput(f"{name}: invalid NAM block dimensions.")
            rn, cn = json.loads(row_names), json.loads(col_names)
            bottom, right = top + len(cn) + nrows - 1, left + len(rn) + ncols - 1
            for old_sheet, t, l, b, r in rectangles:
                if sheet == old_sheet and top <= b and bottom >= t and left <= r and right >= l:
                    raise WrongInput(f"{name}: overlapping NAM data ranges.")
            rectangles.append((sheet, top, left, bottom, right))
            values = sheets[sheet]
            def cell(row, col):
                return values[row - 1][col - 1] if row <= len(values) and col <= len(values[row - 1]) else None
            native_rows = [tuple(cell(top + len(cn) + i, left + k) for k in range(len(rn))) for i in range(nrows)]
            native_cols = [tuple(cell(top + k, left + len(rn) + j) for k in range(len(cn))) for j in range(ncols)]
            data = [[cell(top + len(cn) + i, left + len(rn) + j) for j in range(ncols)] for i in range(nrows)]
            details = json.loads(details)
            details["presence"] = status
            if status == "unavailable":
                if any(v is not None for row in data for v in row):
                    raise WrongInput(f"{name}: unavailable block has values; set presence to observed or remove its values.")
                value = None
            else:
                if any(v is None or isinstance(v, (str, bool)) for row in data for v in row):
                    raise WrongInput(f"{name}: observed blocks require complete numeric values; blanks/formulas are not zeros.")
                value = pd.DataFrame(np.asarray(data, dtype=float).reshape(nrows, ncols),
                                     index=pd.MultiIndex.from_tuples(native_rows, names=rn),
                                     columns=pd.MultiIndex.from_tuples(native_cols, names=cn))
            state.set_block(name, value, scenario, details)
        return state
    finally:
        wb.close()


def write_nam_template(path, *, sets, units, representation="use_imports", block_states=None, item_metadata=None):
    """Create editable blocks, leaving missing observations explicitly unknown."""
    units = dict(units or {})
    valuation = units.pop("valuation", "unspecified")
    schema = NAMSchema(sets, representation=representation, item_metadata=item_metadata or {}, **units)
    state = ModelState(ModelStateMetadata("NAM", name="NAM template", year=int(schema.period) if str(schema.period).isdigit() else None), nam_schema=schema)
    block_states = block_states or {}
    if set(block_states) - set(schema.required_flows()):
        raise WrongInput("Template block_states contains undeclared or inactive flow domains.")
    for name in schema.required_flows():
        status = block_states.get(name, "unavailable")
        if status not in {"unavailable", "declared_zero"}:
            raise WrongInput("An empty NAM template accepts unavailable or declared_zero block states.")
        state.set_block(name, None, metadata={"presence": status, "valuation": valuation})
    return write_nam_workbook(state, path)
