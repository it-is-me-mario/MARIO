"""Local FIGARO-NAM worksheet/bulk normalization with a cell-complete audit."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from scipy import sparse

from mario.internal import ModelState, ModelStateMetadata
from mario.internal.nam import NAMSchema, _coo
from mario.log_exc.exceptions import WrongInput
from mario.model.nam import NAM_BLOCKS, NAM_DOMAINS, NAM_SETS


SOURCE = "Eurostat/JRC FIGARO-NAM local data"
UNSPECIFIED = "Unspecified"
IMPORT_BLOCKS = ("MU", "MY", "Minv", "M_o", "Mcap_d", "Mcap_f")


def _read_rows(path, sheet):
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet not in wb.sheetnames:
                raise WrongInput(f"{path.name}: missing FIGARO-NAM sheet {sheet!r}.")
            return list(wb[sheet].iter_rows(values_only=True))
        finally:
            wb.close()
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, header=None, sep=None, engine="python", dtype=object, keep_default_na=False)
        return [tuple(None if v == "" else v for v in row) for row in frame.itertuples(index=False, name=None)]
    raise WrongInput("FIGARO-NAM requires a local T02 .xlsx workbook or a T02 .csv export with its metadata/header rows.")


def _node(code, section):
    """Interpret source codes only in this source adapter, never in the NAM core."""
    if code.startswith("CPA_"):
        return "commodity", code
    if "i.2 output" in section.lower():
        return "activity", code
    if code in {"D1", "D11", "D12", "B2A3G", "B2", "B3"}:
        return "factor", code
    if code in {"S1", "B9"}:
        return None
    if code == "B8_S1":
        return "other", code
    if re.fullmatch(r"B8_S(11|12|13|14|15|2)", code):
        return ("capital_nonresident" if code == "B8_S2" else "capital_resident"), code[3:]
    if re.fullmatch(r"S(11|12|13|14|15|2)", code):
        return ("income_nonresident" if code == "S2" else "income_resident"), code
    if code in {"D4", "D5", "D61", "D62", "D7", "D8"}:
        return "distribution", code
    if code.startswith("P3_"):
        return "consumption", code
    if code in {"D21X31", "D29X39"}:
        return "tax", code
    if code.startswith("N1"):
        return "nonfinancial_asset", code
    if code == "P51G":
        return "other", code
    if code in {"P52", "P53"}:
        return "nonfinancial_asset", code
    if code == "D9":
        return "distribution", code
    if code == "NP":
        return "nonfinancial_asset", code
    if code == "B9FX9":
        return "other", code
    if re.fullmatch(r"F[1-8][0-9]*", code):
        return "financial_instrument", code
    if code in {"P7", "P33"}:
        return "external_trade", code
    raise WrongInput(f"Unrecognized FIGARO-NAM account {code!r}; normalization cannot discard it.")


def _item_kind(code, group):
    """Keep source subtypes after consolidating logical sets."""
    return {"B8_S1": "investment_allocation", "P51G": "investment_demand",
            "B9FX9": "statistical_adjustment", "P7": "imports", "P33": "consumption_abroad",
            "D9": "capital_transfer", "P52": "inventories", "P53": "valuables",
            "NP": "nonproduced_asset"}.get(code, "fixed_asset" if code.startswith("N1") else group)


@dataclass
class _Ledger:
    path: Path
    region: str
    period: str
    codes: list
    main: np.ndarray
    definitions: dict
    detail_nodes: list
    detail: np.ndarray
    audit: dict

    def account_location(self, code):
        return {key: self.definitions[code][key] for key in ("row", "column")
                if key in self.definitions[code]}

    def cell_location(self, row, column):
        if self.audit["source_format"] == "T02":
            return f"{self.definitions[column]['column']}{self.definitions[row]['row']}"
        return f"base={self.period},ctr={self.region},Set_i={row},m={self.region},Set_j={column}"


def _numeric(path, data, context):
    try:
        values = np.asarray(data, dtype=float)
    except (TypeError, ValueError) as exc:
        raise WrongInput(f"{path.name}: nonnumeric {context} observations.") from exc
    if not np.isfinite(values).all():
        raise WrongInput(f"{path.name}: missing/nonfinite {context} observations; they cannot become zeros.")
    return values


BULK_COLUMNS = ("base", "ctr", "Set_i", "m", "Set_j", "value")


def _is_bulk_csv(path):
    if path.suffix.lower() != ".csv":
        return False
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return tuple(next(csv.reader(stream), ())) == BULK_COLUMNS


@lru_cache(maxsize=1)
def _bulk_accounts():
    # The published sparse release omits zeros, including entire zero accounts.
    # Keep the source classification explicit rather than inferring it from flows.
    with Path(__file__).with_name("figaro_nam_accounts.csv").open(encoding="utf-8-sig", newline="") as stream:
        return {row["Code"]: {"node": None if row["Set"] == "control" else (row["Set"], row["Item"]),
                              "description": row["Name"]} for row in csv.DictReader(stream)}


def _read_bulk_csv(path, *, countries, year=None):
    if isinstance(countries, str):
        selected = [countries]
    else:
        try:
            selected = list(countries) if countries is not None else []
        except TypeError as exc:
            raise WrongInput('FIGARO-NAM bulk CSV requires exactly one country, e.g. countries="IT".') from exc
    if len(selected) != 1 or not isinstance(selected[0], str) or not selected[0].strip():
        raise WrongInput('FIGARO-NAM bulk CSV requires exactly one country, e.g. countries="IT".')
    region = selected[0].strip()
    pieces = []
    with pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", chunksize=100_000) as reader:
        for chunk in reader:
            # m identifies the receiving national ledger; foreign ctr rows are
            # that country's import detail, not extra reference economies.
            selected_rows = chunk.loc[chunk["m"] == region]
            if not selected_rows.empty:
                pieces.append(selected_rows.copy())
    if not pieces:
        raise WrongInput(f"Requested FIGARO-NAM reference economies not present: {[region]}.")
    data = pd.concat(pieces, ignore_index=True)
    for column in BULK_COLUMNS[:-1]:
        if data[column].eq("").any() or data[column].str.strip().ne(data[column]).any():
            raise WrongInput(f"{path.name}: missing or malformed {column} in the selected country.")
    periods = data["base"].unique()
    if len(periods) != 1 or not periods[0].isdigit() or (year is not None and int(periods[0]) != int(year)):
        raise WrongInput(f"{path.name}: expected one reference year matching {year!r}; found {periods.tolist()}.")
    if data.duplicated(list(BULK_COLUMNS[:-1])).any():
        raise WrongInput(f"{path.name}: duplicate FIGARO-NAM bulk observation keys for {region}.")
    data["value"] = _numeric(path, data["value"], "bulk CSV")
    definitions = _bulk_accounts()
    codes = list(definitions)
    unknown = (set(data["Set_i"]) | set(data["Set_j"])) - definitions.keys()
    if unknown:
        raise WrongInput(f"{path.name}: unrecognized FIGARO-NAM bulk account codes: {sorted(unknown)}.")
    national = data.loc[data["ctr"] == region]
    foreign = data.loc[data["ctr"] != region]
    if national.empty:
        raise WrongInput(f"{path.name}: no national ledger observations for {region}.")
    if not foreign["Set_i"].str.startswith("CPA_").all():
        raise WrongInput(f"{path.name}: foreign bulk rows must identify imported products (CPA codes).")
    main = national.pivot(index="Set_i", columns="Set_j", values="value").reindex(index=codes, columns=codes).fillna(0).to_numpy()
    # Expand only observed origin/product pairs, never the all-country Cartesian
    # product. Omitted entries are zeros under this published sparse format.
    if foreign.empty:
        detail_nodes = [(UNSPECIFIED, UNSPECIFIED)]
        detail = np.zeros((1, len(codes)))
    else:
        detail_nodes = list(dict.fromkeys(zip(foreign["ctr"], foreign["Set_i"])))
        detail_index = pd.MultiIndex.from_tuples(detail_nodes, names=("ctr", "Set_i"))
        detail = foreign.pivot(index=["ctr", "Set_i"], columns="Set_j", values="value").reindex(index=detail_index, columns=codes).fillna(0).to_numpy()
    return _Ledger(path, region, periods[0], codes, main, definitions, detail_nodes, detail,
                   {"source_format": "bulk_csv", "reference_country_field": "m",
                    "origin_country_field": "ctr", "selected_observations": len(data),
                    "national_observations": len(national), "bilateral_import_observations": len(foreign),
                    "omitted_observations": "zero within the published FIGARO-NAM sparse classification",
                    "account_classification": "figaro_nam_accounts.csv",
                    "selection": "National ledger and bilateral imports; exports retain national S2 totals.",
                    "consolidated": False})


def _read_t02(path, *, year=None, sheet="T02"):
    rows = _read_rows(path, sheet)
    candidates = [(i, j) for i, row in enumerate(rows) for j, value in enumerate(row)
                  if isinstance(value, str) and "set_i" in value and "set_j" in value]
    if len(candidates) != 1:
        raise WrongInput(f"{path.name}: expected one FIGARO-NAM set_i / set_j header signature.")
    header, code_col = candidates[0]
    first = code_col + 1
    codes = []
    for value in rows[header][first:]:
        if value is None:
            break
        codes.append(str(value).strip())
    if not codes or len(codes) != len(set(codes)) or not {"P7", "S2", "P51G", "B8_S1"} <= set(codes):
        raise WrongInput(f"{path.name}: unsupported FIGARO-NAM account header.")
    n = len(codes)
    main_rows = rows[header + 1:header + 1 + n]
    if len(main_rows) != n or [str(row[code_col]).strip() for row in main_rows] != codes:
        raise WrongInput(f"{path.name}: row/column account codes do not form the same square national ledger.")
    regions = {str(row[code_col - 1]).strip() for row in main_rows}
    periods = {str(row[code_col - 2]).strip() for row in main_rows}
    if len(regions) != 1 or len(periods) != 1:
        raise WrongInput(f"{path.name}: mixed reference economies/periods in one national ledger.")
    region, period = regions.pop(), periods.pop()
    if not period.isdigit() or (year is not None and int(period) != int(year)):
        raise WrongInput(f"{path.name}: reference year {period!r} does not match {year!r}.")
    header_regions = [str(v).strip() for v in rows[header - 1][first:first + n]]
    if header_regions != [region] * n:
        raise WrongInput(f"{path.name}: national columns must retain the reference economy; partner detail belongs below the square.")
    preamble = " ".join(str(v) for row in rows[:header] for v in row if v is not None).casefold()
    if "basic prices" not in preamble or "million euro" not in preamble:
        raise WrongInput(f"{path.name}: T02 basic-price product flows and million-euro units were not identified.")
    main = _numeric(path, [row[first:first + n] for row in main_rows], "national-account")
    definitions, section = {}, ""
    for i, (code, row) in enumerate(zip(codes, main_rows)):
        if row[1] is not None:
            section = str(row[1]).strip()
        definitions[code] = {"node": _node(code, section), "description": row[2],
                             "row": header + 2 + i,
                             "column": get_column_letter(first + i + 1)}
    detail_rows, detail_nodes, detail_values = [], [], []
    for offset, row in enumerate(rows[header + 1 + n:], header + 2 + n):
        if len(row) > code_col and isinstance(row[code_col], str) and row[code_col].startswith("CPA_"):
            if str(row[code_col - 2]).strip() != period:
                raise WrongInput(f"{path.name}:{offset}: mixed periods in bilateral import detail.")
            origin = str(row[code_col - 1]).strip()
            product = row[code_col].strip()
            if not origin or origin == "None":
                raise WrongInput(f"{path.name}:{offset}: missing origin in bilateral import detail.")
            detail_rows.append(offset)
            detail_nodes.append((origin, product))
            detail_values.append(row[first:first + n])
    if not detail_nodes or len(set(detail_nodes)) != len(detail_nodes):
        raise WrongInput(f"{path.name}: missing or duplicate bilateral P7 product-use detail.")
    detail = _numeric(path, detail_values, "bilateral import")
    return _Ledger(path, region, period, codes, main, definitions, detail_nodes, detail,
                   {"source_format": "T02", "sheet": sheet,
                    "bilateral_detail_rows": [min(detail_rows), max(detail_rows)]})


def _normalize_ledger(ledger):
    path, region, period = ledger.path, ledger.region, ledger.period
    codes, main, detail = ledger.codes, ledger.main, ledger.detail
    mapping, memberships = {}, {region: {}}
    controls, source_accounts, item_metadata = [], [], {region: {}}
    for i, code in enumerate(codes):
        definition = ledger.definitions[code]
        node = definition["node"]
        if node is None:
            if np.any(main[i] != 0) or np.any(main[:, i] != 0):
                raise WrongInput(f"{path.name}: nonzero {code} control requires an explicit source normalization rule.")
            controls.append({"code": code, **ledger.account_location(code), "status": "observed_zero_control"})
            continue
        group, item = node
        memberships[region].setdefault(group, []).append(item)
        canonical = (region, NAM_SETS[group]["label"], item)
        if canonical in mapping:
            raise WrongInput(f"{path.name}: duplicate accounting identity after normalization: {canonical}.")
        mapping[canonical] = i
        item_metadata[region].setdefault(group, {})[item] = {
            "kind": _item_kind(code, group), "source_code": code, "description": definition["description"]}
        source_accounts.append({"set": group, "item": item, "code": code,
                                "description": definition["description"], **ledger.account_location(code)})

    user_sets = {"activity", "consumption", "nonfinancial_asset",
                 "capital_resident", "capital_nonresident"}
    supported_columns = {position for node, position in mapping.items()
                         if any(node[1] == NAM_SETS[key]["label"] for key in user_sets)
                         or codes[position] == "P51G"}
    extra_columns = [j for j in range(len(codes)) if j not in supported_columns and np.any(detail[:, j] != 0)]
    if extra_columns:
        raise WrongInput(f"{path.name}: bilateral detail has unsupported nonzero destination accounts: {[codes[j] for j in extra_columns]}.")
    detail_nodes = [(region, origin, product, NAM_SETS["external_trade"]["label"], "P7")
                    for origin, product in ledger.detail_nodes]
    unknown_row = (region, UNSPECIFIED, UNSPECIFIED, NAM_SETS["external_trade"]["label"], "P33")
    if "P33" in codes:
        detail_nodes.append(unknown_row)
    descriptors = {"origin_region": list(dict.fromkeys(node[1] for node in detail_nodes)),
                   "imported_commodity": list(dict.fromkeys(node[2] for node in detail_nodes))}
    layouts = {name: {"rows": ("reference_region", "origin_region", "imported_commodity", "external_trade"),
                      "columns": ("reference_region", NAM_BLOCKS[name]["columns"])} for name in IMPORT_BLOCKS}
    schema = NAMSchema(memberships, representation="use_imports", currency="EUR", scale=1e6,
                       period=int(period), descriptors=descriptors, matrix_layouts=layouts,
                       item_metadata=item_metadata)
    state = ModelState(ModelStateMetadata("NAM", name=f"FIGARO NAM {region} {period}", source=SOURCE,
                                        year=int(period), price="Source valuation: FIGARO-NAM basic-price product flows"), nam_schema=schema)
    # Every nonzero main-square observation must have one declared flow owner.
    owners = {}
    for name in schema.required_flows():
        block = NAM_BLOCKS[name]
        for row_set in NAM_DOMAINS[block["rows"]]["sets"]:
            for col_set in NAM_DOMAINS[block["columns"]]["sets"]:
                owners[NAM_SETS[row_set]["label"], NAM_SETS[col_set]["label"]] = name
    reverse = {position: node for node, position in mapping.items()}
    unowned = []
    for i, j in zip(*np.nonzero(main)):
        if i not in reverse or j not in reverse or (reverse[i][1], reverse[j][1]) not in owners:
            unowned.append(f"{ledger.cell_location(codes[i], codes[j])}: {codes[i]} -> {codes[j]}")
    if unowned:
        raise WrongInput(f"{path.name}: {len(unowned)} nonzero cells lack a NAM flow declaration: {unowned[:8]}.")
    p7 = codes.index("P7")
    differences = main[p7] - detail.sum(axis=0)
    residuals = []
    for j in np.flatnonzero(np.abs(differences) > 1e-6):
        if codes[j] != "S2":
            raise WrongInput(f"{path.name}: unexplained P7/detail difference for {codes[j]}: {differences[j]}.")
        residuals.append({"column_code": codes[j], "value": float(differences[j]),
                          "native_block": "TP_trade", "meaning": "Observed non-product external settlement; no bilateral allocation inferred.",
                          "cell": ledger.cell_location("P7", codes[j])})
    imported_index = pd.MultiIndex.from_tuples(detail_nodes, names=("Region", "origin_region", "imported_commodity", "Set", "Item"))
    for name in schema.required_flows():
        block = NAM_BLOCKS[name]
        columns = schema.accounts(block["columns"])
        col_positions = [mapping[node] for node in columns]
        if name in IMPORT_BLOCKS:
            values = detail[:, col_positions]
            if "P33" in codes:
                values = np.vstack([values, main[codes.index("P33"), col_positions]])
            index = imported_index
        else:
            index = schema.accounts(block["rows"])
            values = main[np.ix_([mapping[node] for node in index], col_positions)]
        frame = pd.DataFrame.sparse.from_spmatrix(sparse.csr_matrix(values), index=index, columns=columns)
        state.set_block(name, frame, metadata={"presence": "observed", "valuation": "source_T02" if ledger.audit["source_format"] == "T02" else "source_FIGARO_NAM",
                                              "source_file": path.name, "source_format": ledger.audit["source_format"],
                                              **({"source_sheet": ledger.audit["sheet"]} if "sheet" in ledger.audit else {}),
                                              "source_component": "bilateral P7 plus unallocated P33" if name in IMPORT_BLOCKS else "national square"})
    nodes = list(mapping)
    positions = list(mapping.values())
    reconstructed = state.compute("T").reindex(index=pd.MultiIndex.from_tuples(nodes), columns=pd.MultiIndex.from_tuples(nodes)).sparse.to_dense().to_numpy()
    expected = main[np.ix_(positions, positions)]
    max_error = float(np.max(np.abs(reconstructed - expected)))
    if not np.allclose(reconstructed, expected, atol=1e-6, rtol=1e-10):
        raise WrongInput(f"{path.name}: native blocks fail to reproduce the source national square (max error {max_error}).")
    state.scenarios["baseline"].blocks.pop("T", None)
    state.metadata.extra["figaro_nam"] = {
        "source_files": [path.name], **ledger.audit, "reference_economies": [region],
        "source_accounts": source_accounts, "controls": controls,
        "external_settlements": residuals, "source_reconciled": True,
        "max_cell_reconstruction_error": max_error,
        "max_source_account_imbalance": float(np.max(np.abs(main.sum(axis=1) - main.sum(axis=0)))),
        "financial_reconciliation": "B9_d/B9_f are financial-side only; B9FX9 stays in Other accounts (O_sc_d/O_sc_f).",
        "unspecified_descriptor": {"value": UNSPECIFIED, "meaning": "No origin/product observed for P33; not a geographic region."},
    }
    state.metadata.add_history(f"Normalized {path.name} ({ledger.audit['source_format']}, {region}); all nonzero source cells assigned and reconstructed.")
    return state


def parse_figaro_nam(path, *, year=None, countries=None, sheet="T02"):
    """Read one or more harmonized national ledgers without consolidating them."""
    if path is None:
        raise WrongInput("FIGARO-NAM needs path= to a local workbook; NAM is not requested through the IOT/SUT API cache.")
    if isinstance(path, (str, Path)):
        root = Path(path)
        paths = sorted(root.glob("FIGARO_NAM_*.*")) if root.is_dir() else [root]
    else:
        paths = [Path(value) for value in path]
    paths = [p for p in paths if p.suffix.lower() in {".xlsx", ".xlsm", ".csv"}]
    if not paths:
        raise WrongInput("No local FIGARO-NAM workbooks/T02 CSV exports found.")
    bulk_paths = [file for file in paths if _is_bulk_csv(file)]
    if bulk_paths:
        if len(paths) != 1:
            raise WrongInput("Pass one FIGARO-NAM bulk CSV path at a time; do not mix it with national workbooks or other years.")
        return _normalize_ledger(_read_bulk_csv(bulk_paths[0], countries=countries, year=year))
    wanted = None if countries is None else ({countries} if isinstance(countries, str) else set(countries))
    states, seen = [], set()
    for file in paths:
        current = _normalize_ledger(_read_t02(file, year=year, sheet=sheet))
        region = next(iter(current.nam_schema.memberships))
        if wanted is not None and region not in wanted:
            continue
        if region in seen:
            raise WrongInput(f"Duplicate FIGARO-NAM reference economy {region}; select one source file per economy.")
        states.append(current)
        seen.add(region)
    if not states or (wanted is not None and wanted - seen):
        raise WrongInput(f"Requested FIGARO-NAM reference economies not present: {sorted((wanted or set()) - seen)}.")
    if len({s.nam_schema.period for s in states}) != 1:
        raise WrongInput("A multi-country NAM requires a common reference period.")
    if len(states) == 1:
        return states[0]
    first = states[0]
    members = {region: dict(groups) for state in states for region, groups in state.nam_schema.memberships.items()}
    descriptors = {role: list(dict.fromkeys(item for state in states for item in state.nam_schema.descriptors[role]))
                   for role in first.nam_schema.descriptors}
    schema = NAMSchema(members, representation="use_imports", currency="EUR", scale=1e6,
                       period=first.nam_schema.period, descriptors=descriptors,
                       matrix_layouts={name: dict(layout) for name, layout in first.nam_schema.matrix_layouts.items()},
                       item_metadata={r: groups for state in states
                                      for r, groups in state.nam_schema.to_dict()["item_metadata"].items()})
    metadata = ModelStateMetadata("NAM", name=f"FIGARO NAM {schema.period} ({len(states)} economies)", source=SOURCE,
                                  year=first.metadata.year, price=first.metadata.price,
                                  extra={"figaro_nam": {"sources": [state.metadata.extra["figaro_nam"] for state in states],
                                                        "reference_economies": list(members), "source_reconciled": True,
                                                        "consolidated": False}})
    result = ModelState(metadata, nam_schema=schema)
    for name in schema.required_flows():
        frames = [state.get_block(name) for state in states if state.has_matrix(name)]
        if not frames:
            continue
        rows = frames[0].index.append([frame.index for frame in frames[1:]])
        columns = schema.accounts(NAM_BLOCKS[name]["columns"])
        rr, cc, vv = [], [], []
        for frame in frames:
            coo = _coo(frame)
            rr.extend(rows.get_indexer(frame.index)[coo.row])
            cc.extend(columns.get_indexer(frame.columns)[coo.col])
            vv.extend(coo.data)
        frame = pd.DataFrame.sparse.from_spmatrix(sparse.csr_matrix((vv, (rr, cc)), shape=(len(rows), len(columns))), index=rows, columns=columns)
        result.set_block(name, frame, metadata={"presence": "observed", "valuation": "source_T02", "source_files": [p.name for p in paths]})
    return result
