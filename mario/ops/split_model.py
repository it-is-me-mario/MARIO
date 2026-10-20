"""Opt-in convex IOT split experiments and explicit inherited accounting.

The supplied baseline is never closed or edited. With accounting='inherit',
each preceding row/column discrepancy is distributed within its split family
by fixed output shares, including the residual parent. A zero-output family
retains its discrepancy in the parent. No discrepancy is a decision variable.
"""
from __future__ import annotations

import json
from pathlib import Path
import time

import numpy as np
import pandas as pd

from mario.log_exc.exceptions import WrongInput


DEFAULTS = dict(accounting='strict', value_added_loss='legacy', flow_loss='kl', final_demand_weight=0.,
                supply_weight=0., complete_intermediate_priors=False,
                check_feasibility=False, feasibility_solver='HIGHS', scale=1.,compact_constraints=False,
                defer_satellite_recalculation=False)
FLOW_TABLES = {'Z0','V0','X','Trade','Zold','Yold','Vold','Z_refine',
               'RowResidual','ColumnResidual','Y0','Z_use0','Z_supply0'}


def normalize_options(options):
    if options is None:
        return None
    if not isinstance(options, dict) or set(options)-DEFAULTS.keys():
        raise WrongInput(f'split_model_options must be a dict with keys {sorted(DEFAULTS)}.')
    result = {**DEFAULTS, **options}
    if result['accounting'] not in ('strict','inherit'):
        raise WrongInput("accounting must be 'strict' or 'inherit'.")
    if result['value_added_loss'] not in ('legacy','signed_kl','quadratic'):
        raise WrongInput("value_added_loss must be 'legacy', 'signed_kl' or 'quadratic'.")
    if result['flow_loss'] not in ('kl','quadratic'):
        raise WrongInput("flow_loss must be 'kl' or 'quadratic'.")
    if result['flow_loss']=='quadratic' and not result['complete_intermediate_priors']:
        raise WrongInput('Quadratic flow loss requires complete_intermediate_priors=True.')
    for key in ('final_demand_weight','supply_weight','scale'):
        value = result[key]
        if isinstance(value, bool) or not isinstance(value, (int,float)) or not np.isfinite(value) or value < 0 or (key=='scale' and value==0):
            raise WrongInput(f'{key} must be finite and {"positive" if key=="scale" else "nonnegative"}.')
    for key in ('complete_intermediate_priors','check_feasibility','compact_constraints',
                'defer_satellite_recalculation'):
        if not isinstance(result[key], bool):
            raise WrongInput(f'{key} must be boolean.')
    if not isinstance(result['feasibility_solver'],str) or not result['feasibility_solver']:
        raise WrongInput('feasibility_solver must name an installed LP solver.')
    return result


def inherited_accounting(before, after_x, parents, scenario='baseline'):
    """Return row/column discrepancy targets aligned with the refined X index."""
    z, y, v, x = (before.resolve(name, scenario=scenario) for name in ('Z','Y','V','X'))
    original_x = x.iloc[:,0]
    original_row = z.sum(axis=1)+y.sum(axis=1)-original_x.reindex(z.index)
    original_col = z.sum(axis=0)+v.sum(axis=0)-original_x.reindex(z.columns)
    quantities = after_x.iloc[:,0]
    row, col = [], []
    for label, output in quantities.items():
        source = (*label[:-1], parents.get(label[-1],label[-1]))
        if source not in original_x.index:
            raise WrongInput(f'No preceding output/accounting reference for {label!r}.')
        denominator = float(original_x.loc[source])
        if denominator < 0 or output < -1e-10:
            raise WrongInput('Inherited accounting requires nonnegative outputs.')
        share = float(output)/denominator if denominator>0 else float(label==source)
        row.append(share*float(original_row.loc[source]))
        col.append(share*float(original_col.loc[source]))
    return pd.DataFrame({'row':row,'column':col}, index=quantities.index)


def _values(prior):
    value = prior.value
    if value is None:
        raise ValueError('A split prior must have a value before model construction.')
    if hasattr(value,'toarray'):
        value = value.toarray()
    result = np.asarray(value,dtype=float)
    if not np.isfinite(result).all():
        raise ValueError('Nonfinite split prior.')
    return result


def _array(value):
    return np.asarray(value.toarray() if hasattr(value,'toarray') else value,dtype=float)


def model_active_kl(a, prior, zero_mask=None, signed=False, reference_floor=1e-8):
    """KL on allowed cells; structural zeros create no exponential cones."""
    import cvxpy as cp
    q = _values(prior)
    allowed = np.ones(q.size,dtype=bool) if zero_mask is None else _values(zero_mask).reshape(-1)==0
    if not allowed.any():
        return cp.Constant(0.)
    q = q.reshape(-1)[allowed]
    if not signed and (q<0).any():
        raise ValueError('Intermediate-input priors must be nonnegative.')
    vector = cp.reshape(a,(a.size,),order='C')[np.flatnonzero(allowed)]
    if signed:
        vector = cp.multiply(np.where(q<0,-1.,1.),vector)
    # The floor is numerical support, not a source observation. Structural
    # exclusions remain fixed independently of the objective.
    return cp.sum(cp.kl_div(vector,np.abs(q)+reference_floor))


def model_signed_kl(a, prior, zero_mask=None, reference_floor=1e-8):
    return model_active_kl(a,prior,zero_mask,signed=True,reference_floor=reference_floor)


def model_quadratic(a, prior, zero_mask=None, reference_floor=1e-8):
    import cvxpy as cp
    q = _values(prior)
    allowed = np.ones(q.size,dtype=bool) if zero_mask is None else _values(zero_mask).reshape(-1)==0
    if not allowed.any():
        return cp.Constant(0.)
    vector = cp.reshape(a,(a.size,),order='C')[np.flatnonzero(allowed)]
    q = q.reshape(-1)[allowed]
    scale = np.maximum(np.abs(q),max(float(np.max(np.abs(q),initial=0.))*1e-6,reference_floor))
    return cp.sum_squares(cp.multiply(1/np.sqrt(scale),vector-q))


def model_prior_signed_values(a, prior, zero_mask=None):
    """Explicit prior-sign domain, also available without entropy atoms."""
    import cvxpy as cp
    q = _values(prior).reshape(-1)
    allowed = np.ones(q.size,dtype=bool) if zero_mask is None else _values(zero_mask).reshape(-1)==0
    if not allowed.any():
        return cp.Constant(0.)
    vector = cp.reshape(a,(a.size,),order='C')[np.flatnonzero(allowed)]
    return cp.multiply(np.where(q[allowed]<0,-1.,1.),vector)


def model_legacy_entropy(a, prior, *, signed=False, reference_floor=1e-8):
    """Original matrix-valued loss, with a scale-covariant numerical offset."""
    import cvxpy as cp
    q = _values(prior)
    return cp.rel_entr(cp.multiply(np.where(q<0,-1.,1.),a) if signed else a,
                       np.abs(q)+reference_floor if signed else prior+reference_floor)


def model_masked_values(a, mask):
    """Keep constrained coordinates only; omit identities such as 0 == 0."""
    import cvxpy as cp
    positions = np.flatnonzero(_values(mask).reshape(-1)!=0)
    return cp.reshape(a,(a.size,),order='C')[positions] if positions.size else cp.Constant(0.)


def proportional_row_reference(table, original, outputs, previous_outputs, parents):
    """Joint output-share prior for Y or deliveries to unchanged buyers.

    Always read the preceding parent row, never a row already reduced by an
    earlier child. This makes the reference independent of child ordering and
    conserves each preceding cell across the child/residual-parent family.
    """
    keys = [column for column in table.columns if column not in ('id','values')]
    source = table[keys].copy()
    source['sector_from_Name'] = source['sector_from_Name'].map(lambda name:parents.get(name,name))
    values = original.set_index(keys)['values'].reindex(pd.MultiIndex.from_frame(source)).to_numpy()

    def output_series(frame):
        return pd.Series(frame.iloc[:,0].to_numpy(),index=pd.MultiIndex.from_tuples(
            [(label[0],label[-1]) for label in frame.index]))

    target_index = pd.MultiIndex.from_frame(table[['region_from_Name','sector_from_Name']])
    source_index = pd.MultiIndex.from_frame(source[['region_from_Name','sector_from_Name']])
    numerator = output_series(outputs).reindex(target_index).to_numpy(dtype=float)
    denominator = output_series(previous_outputs).reindex(source_index).to_numpy(dtype=float)
    share = np.divide(numerator,denominator,
                      out=table['sector_from_Name'].eq(source['sector_from_Name']).to_numpy(dtype=float),
                      where=denominator>0)
    if (not np.isfinite(values).all() or not np.isfinite(numerator).all()
            or not np.isfinite(denominator).all() or (numerator<0).any() or (denominator<0).any()):
        raise WrongInput('Incomplete or invalid preceding row/output reference for split priors.')
    return table.assign(values=values*share)


def configure_model(path, options):
    """Extend only the generated CVXLab configuration, never its template."""
    if options is None:
        return
    from openpyxl import load_workbook
    path = Path(path)
    book = load_workbook(path)
    ws = book['structure_variables']
    headers = [c.value for c in ws[1]]
    layouts = {row[headers.index('variables_info')].value: [c.value for c in row]
               for row in ws.iter_rows(min_row=2)}

    def add(source, name, *, table=None, kind=None, **coordinates):
        row = list(layouts[source])
        changes = dict(variables_info=name,table_key=table or name,
                       description='Explicit experimental IOT split reference')
        if kind:
            changes.update(type=kind,nonneg=None)
        changes.update(coordinates)
        for key,value in changes.items():
            if key in headers:
                row[headers.index(key)] = value
        ws.append(row)

    if options['accounting']=='inherit':
        add('X_intra','RowResidual',kind='exogenous')
        add('XT_intra','ColumnResidual',kind='exogenous')
    if options['final_demand_weight']:
        add('Y_pn_to','Y_all',table='Y',region_to='dim: cols')
        add('Y_pn_to','Y0',kind='exogenous',region_to='dim: cols')
        add('Zero_Y_mask','Zero_Y_all',table='Zero_Y_mask',region_to='dim: cols')
    if options['complete_intermediate_priors']:
        add('Z_use_to','Z_use_all',table='Z_use',region_to='dim: cols')
        add('Z_use_to','Z_use0',kind='exogenous',region_to='dim: cols')
        add('Zero_use_mask','Zero_use_all',table='Zero_use_mask',region_to='dim: cols')
    if options['supply_weight'] or options['complete_intermediate_priors']:
        add('Z_supply_to','Z_supply_all',table='Z_supply',region_to='dim: cols')
        add('Z_supply_to','Z_supply0',kind='exogenous',region_to='dim: cols')
        add('Zero_supply_mask','Zero_supply_all',table='Zero_supply_mask',region_to='dim: cols')

    problem = book['problem']
    cols = {c.value:c.column for c in problem[1]}
    flow_loss = 'model_quadratic' if options['flow_loss']=='quadratic' else 'model_active_kl'
    y_loss = 'model_quadratic' if options['flow_loss']=='quadratic' else 'model_signed_kl'
    for row in problem.iter_rows(min_row=2):
        objective = row[cols['objective']-1]
        if objective.value:
            if options['complete_intermediate_priors']:
                terms = [f'{flow_loss}(Z_use_all,Z_use0,Zero_use_all)']
            else:
                terms = ['sum(entropy(Z_use_n,Z0_all_n))','sum(entropy(Z_use_st_p,Z0_st_p))']
            loss = options['value_added_loss']
            terms.append({'legacy':'sum(entropy_signed(V,V0))',
                          'signed_kl':'model_signed_kl(V,V0,Zero_V_mask)',
                          'quadratic':'model_quadratic(V,V0,Zero_V_mask)'}[loss])
            if options['final_demand_weight']:
                terms.append(f"{options['final_demand_weight']}*{y_loss}(Y_all,Y0,Zero_Y_all)")
            if options['supply_weight'] or options['complete_intermediate_priors']:
                weight = options['supply_weight'] or 1.
                terms.append(f'{weight}*{flow_loss}(Z_supply_all,Z_supply0,Zero_supply_all)')
            objective.value = 'Minimize('+'+'.join(terms)+')'
        expression = row[cols['expressions']-1]
        value = expression.value or ''
        if options['accounting']=='inherit':
            if value.startswith('X_intra'):
                expression.value = ('Z_supply_from@I_r_st+Z_use_pn_pn_from@I_r_pn+Y_pn_from@I_r_c'
                                    ' == X_intra+RowResidual') if '<=' in value else None
            elif value.startswith('XT_intra'):
                expression.value = ('I_row_rs@Z_use_to+I_row_v@V_intra == XT_intra+ColumnResidual') if '<=' in value else None
    explicit_domains = []
    if options['value_added_loss']=='quadratic':
        explicit_domains.append('model_prior_signed_values(V,V0,Zero_V_mask) >= 0')
    if options['flow_loss']=='quadratic' and options['final_demand_weight']:
        explicit_domains.append('model_prior_signed_values(Y_all,Y0,Zero_Y_all) >= 0')
    for expression in explicit_domains:
        row = [None]*problem.max_column
        row[cols['expressions']-1] = expression
        problem.append(row)
    book.save(path)
    book.close()
    with (path.parent/'user_defined_operators.py').open('a') as stream:
        for name in ('model_active_kl','model_signed_kl','model_quadratic'):
            stream.write(f'\n\ndef {name}(*args):\n    from mario.ops.split_model import {name} as operation\n    return operation(*args, reference_floor={1e-8/options["scale"]!r})\n')
        if options['scale']!=1:
            for name,signed in [('entropy',False),('entropy_signed',True)]:
                stream.write(f'\n\ndef {name}(a,prior):\n    from mario.ops.split_model import model_legacy_entropy\n    return model_legacy_entropy(a,prior,signed={signed!r},reference_floor={1e-8/options["scale"]!r})\n')
        stream.write('\n\ndef model_masked_values(*args):\n    from mario.ops.split_model import model_masked_values as operation\n    return operation(*args)\n')
        stream.write('\n\ndef model_prior_signed_values(*args):\n    from mario.ops.split_model import model_prior_signed_values as operation\n    return operation(*args)\n')
    (path.parent/'split_model_options.json').write_text(json.dumps(options,indent=2,allow_nan=False))


def compact_model(path, options, *, exact_trade=False):
    if not options or not options['compact_constraints']:
        return
    import re
    from openpyxl import load_workbook
    book = load_workbook(path)
    sheet = book['problem']
    column = [c.value for c in sheet[1]].index('expressions')+1
    for row in sheet.iter_rows(min_row=2):
        cell = row[column-1]
        value = cell.value or ''
        zero = re.fullmatch(r'\s*0\s*==\s*mult\((Zero_\w+),\s*(\w+)\)\s*',value)
        sign = re.fullmatch(r'mult\(((?:Positive|Negative)_Y_mask),\s*(\w+)\)\s*([<>]=)\s*0',value)
        if zero:
            cell.value = f'model_masked_values({zero[2]},{zero[1]}) == 0'
        elif sign:
            cell.value = f'model_masked_values({sign[2]},{sign[1]}) {sign[3]} 0'
        elif options['accounting']=='inherit':
            if value.startswith('Zold_') and '<=' in value and value.endswith('*(1+eps)'):
                left,right = value.split('<=',1)
                cell.value = left+' == '+right[:-len('*(1+eps)')]
            elif value.startswith('Zold_') and '>=' in value and value.endswith('*(1-eps)'):
                cell.value = None
            for name in ('Yold','Vold'):
                if value.startswith(f'{name}-absval({name})*eps<='):
                    cell.value = name+' == '+value.split('<=',1)[1]
                elif value.startswith(f'{name}+absval({name})*eps>='):
                    cell.value = None
        if exact_trade and value.startswith('Trade'):
            if '<=' in value and value.endswith('*(1+delta)'):
                left,right = value.split('<=',1)
                cell.value = left+' == '+right[:-len('*(1+delta)')]
            elif '>=' in value and value.endswith('*(1-delta)'):
                cell.value = None
    book.save(path)
    book.close()


def fill_model_inputs(instance, input_data, flat_matrices, set_map, scenario_label, options):
    if options is None:
        return
    if options['accounting']=='inherit' and (input_data['tol'].loc[input_data['tol']['scalar_Name'].eq('eps'),'values']!=0).any():
        raise WrongInput('Inherited accounting requires exact preceding aggregates (eps=0).')
    from mario.ops.cvxlab_bridge import _rename_flat_columns, _split_parent_lookup, _merge_cvxlab_input_table
    split = f'split_{scenario_label}'
    for name,matrix in [('Y0','Y'),('Z_use0','Z'),('Z_supply0','Z')]:
        if name in input_data:
            if name=='Z_use0':
                values = flat_matrices[matrix].query('Scenario == @split').drop(columns=['Scenario'])
                input_data[name] = _merge_cvxlab_input_table(input_data[name],_rename_flat_columns(values,set_map),table_name=name)
            else:
                original = flat_matrices[matrix].query("Scenario == 'original'").drop(columns=['Scenario'])
                input_data[name] = proportional_row_reference(
                    input_data[name],_rename_flat_columns(original,set_map),instance.matrices[split]['X'],
                    instance.resolve('X',scenario='original'),_split_parent_lookup(instance))
    refinement = getattr(instance,'_intermediate_refinement',None)
    if 'Z_supply0' in input_data and refinement is not None and not refinement.empty:
        keys = ['region_from_Name','sector_from_Name','region_to_Name','sector_to_Name']
        current = input_data['Z_supply0'].set_index(keys)
        updates = refinement.set_index(keys)['values']
        hit = current.index.intersection(updates.index)
        current.loc[hit,'values'] = updates.reindex(hit)
        input_data['Z_supply0'] = current.reset_index()
    if options['accounting']=='inherit':
        target = inherited_accounting(instance,instance.matrices[split]['X'],_split_parent_lookup(instance),scenario='original')
        for table,column,keys in [('RowResidual','row',['region_from_Name','sector_from_Name']),
                                  ('ColumnResidual','column',['region_to_Name','sector_to_Name'])]:
            series = pd.Series(target[column].to_numpy(),index=pd.MultiIndex.from_tuples(
                [(label[0],label[-1]) for label in target.index],names=keys))
            index = pd.MultiIndex.from_frame(input_data[table][keys])
            input_data[table]['values'] = series.reindex(index).to_numpy()
    for name in ('RowResidual','ColumnResidual','Y0','Z_use0','Z_supply0'):
        if name in input_data and not np.isfinite(input_data[name]['values'].to_numpy(dtype=float)).all():
            raise WrongInput(f'Incomplete or nonfinite {name} model inputs.')


def check_feasibility(model, destination, options):
    if not options or not options['check_feasibility']:
        return
    import cvxpy as cp
    solver = options['feasibility_solver']
    if solver not in cp.installed_solvers():
        raise WrongInput(f'Feasibility solver {solver} is not installed.')
    records = []
    try:
        for key,frame in model.core.problem.numerical_problems.items():
            for scenario,row in frame.iterrows():
                original = row['problem']
                # Objective atoms can impose sign/domain constraints absent
                # from Problem.constraints. Dropping those changes feasibility.
                constraints = original.constraints+original.objective.expr.domain
                if not all(c.expr.is_affine() for c in constraints):
                    raise WrongInput('The feasibility stage requires affine split constraints/domains.')
                feasible = cp.Problem(cp.Minimize(0),constraints)
                started = time.monotonic()
                kwargs = {'highs_options':{'threads':1,'primal_feasibility_tolerance':1e-9,
                                          'dual_feasibility_tolerance':1e-9}} if solver=='HIGHS' else {}
                feasible.solve(solver=solver,ignore_dpp=True,**kwargs)
                record = dict(problem=str(key),scenario=str(scenario),status=feasible.status,
                              seconds=time.monotonic()-started,variables=sum(v.size for v in feasible.variables()),
                              constraints=sum(c.size for c in constraints))
                if feasible.status in ('optimal','optimal_inaccurate'):
                    record['max_constraint_violation_scaled'] = max(
                        (float(_array(c.violation()).max(initial=0.)) for c in constraints),default=0.)
                records.append(record)
                if feasible.status not in ('optimal','optimal_inaccurate'):
                    raise WrongInput(f'Linear split feasibility check returned {feasible.status}; see feasibility.json.')
    finally:
        (Path(destination)/'feasibility.json').write_text(json.dumps(records,indent=2,allow_nan=False))


def verify_model_solution(model, destination, options):
    """Report and validate every affine constraint/domain in original units."""
    if options is None:
        return
    records = []
    failures = 0
    for key,frame in model.core.problem.numerical_problems.items():
        for scenario,row in frame.iterrows():
            problem = row['problem']
            bad, worst = 0, 0.
            for constraint in problem.constraints+problem.objective.expr.domain:
                violation = _array(constraint.violation())
                if not np.isfinite(violation).all():
                    raise WrongInput('Nonfinite model solution constraint residual.')
                reference = np.maximum(np.abs(_array(constraint.args[0].value)),np.abs(_array(constraint.args[1].value)))
                bad += int(np.count_nonzero(violation > 1e-6/options['scale']+1e-8*reference))
                worst = max(worst,float(violation.max(initial=0.))*options['scale'])
            stats = problem.solver_stats
            objective = float(problem.value)*options['scale']
            records.append(dict(problem=str(key),scenario=str(scenario),status=problem.status,
                                objective=objective if np.isfinite(objective) else None,
                                objective_finite=bool(np.isfinite(objective)),failed_constraint_cells=bad,
                                max_constraint_violation=worst,solver_seconds=stats.solve_time))
            if not np.isfinite(objective):
                records[-1]['warning'] = ('Objective evaluation is nonfinite at the returned point; '
                                          'all explicit and implicit domain residuals are checked independently.')
            failures += bad
    (Path(destination)/'model_solution.json').write_text(json.dumps(records,indent=2,allow_nan=False))
    if failures:
        raise WrongInput(f'{failures} split constraint cells exceed numerical precision; see model_solution.json.')
