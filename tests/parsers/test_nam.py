"""Portable FIGARO-NAM source and native-workbook tests (no network/files)."""

import csv

import numpy as np
import pandas.testing as pdt
import pytest
from openpyxl import Workbook, load_workbook

import mario
from mario.compute.resolver import ResolutionError
from mario.internal import ModelState
from mario.log_exc.exceptions import DataMissing, NotImplementable, WrongInput
from mario.model.nam import NAM_SETS, NAM_BLOCKS
from mario.parsers.nam_workbook import read_nam_workbook


def write_figaro_source(path, region="IT", year=2023, header=12, change=None):
    codes = ["CPA_A01", "CPA_C10", "A01", "C10", "D11", "S1", "S14", "S2",
             "P3_S14", "D21X31", "D29X39", "B8_S1", "B8_S14", "B8_S2",
             "N111G", "N1131G", "P51G", "D9", "NP", "B9", "F21", "F41", "B9FX9", "P33", "P7"]
    x = np.zeros((len(codes), len(codes)))
    def pay(row, col, value):
        x[codes.index(row), codes.index(col)] = value
    for row, col, value in [
        ("A01", "CPA_A01", 100), ("CPA_A01", "A01", 40),
        ("D11", "A01", 50), ("D11", "S2", 3),
        ("S14", "D11", 51), ("S2", "D11", 2),
        ("P3_S14", "S14", 46), ("CPA_A01", "P3_S14", 40),
        ("CPA_A01", "P51G", 10), ("CPA_A01", "S2", 10),
        ("D21X31", "P51G", 2), ("S14", "D21X31", 2),
        ("B8_S14", "S14", 7), ("B8_S2", "S2", 7),
        ("N111G", "B8_S14", 8), ("N1131G", "B8_S14", 6),
        ("B8_S1", "N111G", 8), ("B8_S1", "N1131G", 6), ("P51G", "B8_S1", 14),
        ("P7", "A01", 10), ("P7", "P3_S14", 5), ("P7", "P51G", 2),
        ("P7", "S2", 2), ("S2", "P7", 19), ("P33", "P3_S14", 1), ("S2", "P33", 1),
        ("F41", "B8_S2", 8), ("B8_S14", "F41", 8),
        ("B9FX9", "B8_S14", 1), ("B9FX9", "B8_S2", -1),
    ]:
        pay(row, col, value)
    if change:
        change(codes, x)
    rows = [[None] * (len(codes) + 7) for _ in range(header - 1)]
    rows[1][1:3] = ["Table 2", "FIGARO - National Accounting Matrix (FNAM)"]
    rows[2][1:3] = ["Country:", region]
    rows[3][1:3] = ["Year:", str(year)]
    rows[4][1:3] = ["Unit:", "Current prices, million euro"]
    rows[7][1] = "F-NAM (Basic Prices)"
    rows[header - 2][6:] = ["counterpart_area", *([region] * len(codes))]
    rows.append([None] * 6 + ["set_i / set_j", *codes])
    for i, code in enumerate(codes):
        section = "I.1 Goods and services" if i == 0 else ("I.2 Output" if i == 2 else ("II.1 Earned income" if i == 4 else None))
        rows.append([None, section, f"Description of {code}", i + 1, str(year), region, code, *x[i]])
    for origin, product, amounts in [("DE", "CPA_A01", {"A01": 4, "P51G": 2}),
                                      ("NL", "CPA_C10", {"A01": 6, "P3_S14": 5})]:
        values = [amounts.get(code, 0) for code in codes]
        rows.append([None, None, f"Imported {product}", len(rows), year, origin, product, *values])
    if path.suffix == ".csv":
        with path.open("w", newline="") as f:
            csv.writer(f).writerows(rows)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "T02"
        for row in rows:
            ws.append(row)
        wb.save(path)
    return codes, x


def dense(frame):
    return frame.sparse.to_dense() if all(str(t).startswith("Sparse") for t in frame.dtypes) else frame


def test_public_figaro_nam_preserves_every_source_flow_and_native_import_detail(tmp_path):
    source = tmp_path / "FIGARO_NAM_IT_2023.xlsx"
    write_figaro_source(source, header=15)
    db = mario.parse_figaro(source, table="NAM")
    assert isinstance(db, mario.Database)
    assert db.table_type == "NAM"
    assert db.get_index("Region") == ["IT"]
    assert db.get_index("capital_resident") == db.get_index("income_resident") == ["S14"]
    assert db.get_index("o") == ["B8_S1", "P51G", "B9FX9"]
    assert db.get_index("tr") == ["P33", "P7"]
    assert dense(db.O_o).loc[("IT", "Other accounts", "P51G"), ("IT", "Other accounts", "B8_S1")] == 14
    assert dense(db.O_nf).to_numpy().sum() == 14
    assert db.MU.index.names == ["Region", "origin_region", "imported_commodity", "Set", "Item"]
    assert dense(db.MU).iloc[0, 0] == 4
    assert dense(db.MU).iloc[1, 0] == 6
    assert "Unspecified" in db.MY.index.get_level_values("origin_region")
    assert dense(db.MY).iloc[-1, 0] == 1
    assert dense(db.TP_trade).to_numpy().sum() == 2
    assert db.info["figaro_nam"]["source_reconciled"]
    assert db.info["figaro_nam"]["max_cell_reconstruction_error"] < 1e-8
    assert db.info["figaro_nam"]["external_settlements"][0]["native_block"] == "TP_trade"
    view = dense(db.T)
    np.testing.assert_allclose(view.sum(axis=0), view.sum(axis=1))
    capital = ("IT", NAM_SETS["capital_resident"]["label"], "S14")
    assert db.B9_d.loc["B9", capital] == -8
    assert dense(db.O_sc_d).iloc[-1, 0] == 1
    for operation in ("aggregate", "calc_linkages", "is_isard"):
        with pytest.raises(NotImplementable, match="NAM"):
            getattr(db, operation)


@pytest.mark.parametrize("extension", ["xlsx", "csv"])
def test_figaro_local_formats_infer_year_and_never_request_api(tmp_path, monkeypatch, extension):
    source = tmp_path / f"FIGARO_NAM_IT_2023.{extension}"
    write_figaro_source(source)
    monkeypatch.setattr("mario.parsers.figaro._download_figaro_api_frame", lambda *a, **k: pytest.fail("NAM must not request IOT/SUT API"))
    state = mario.parse_figaro(source, table="NAM", model="ModelState")
    assert isinstance(state, ModelState)
    assert state.metadata.year == 2023
    with pytest.raises(WrongInput, match="year"):
        mario.parse_figaro(source, table="NAM", year=2022)
    with pytest.raises(WrongInput, match="reference economies"):
        mario.parse_figaro(source, table="NAM", countries=["FR"])


def test_multi_country_files_preserve_reference_economy_and_partner_roles(tmp_path):
    files = [tmp_path / f"FIGARO_NAM_{region}_2023.xlsx" for region in ("IT", "FR")]
    for path, region in zip(files, ("IT", "FR")):
        write_figaro_source(path, region=region)
    db = mario.parse_figaro(files, table="NAM")
    assert db.get_index("Region") == ["IT", "FR"]
    assert db.info["figaro_nam"]["consolidated"] is False
    use = dense(db.U)
    assert use.loc[("FR", "Commodity", "CPA_A01"), ("FR", "Activity", "A01")] == 40
    assert use.loc[("FR", "Commodity", "CPA_A01"), ("IT", "Activity", "A01")] == 0
    assert set(db.MU.index.get_level_values("Region")) == {"IT", "FR"}
    assert set(db.MU.index.get_level_values("origin_region")) == {"DE", "NL", "Unspecified"}
    assert "Region = 2\n" in repr(db)
    assert "Commodity = 2\n" in repr(db)
    selected = mario.parse_figaro(tmp_path, table="NAM", countries="FR")
    assert selected.get_index("Region") == ["FR"]
    output = tmp_path / "native_multi_country.xlsx"
    db.to_excel(output)
    restored = mario.parse_from_excel(output, table="NAM", mode="flows")
    pdt.assert_frame_equal(dense(db.T), dense(restored.T), check_dtype=False)
    with pytest.raises(WrongInput, match="Duplicate"):
        mario.parse_figaro([files[0], files[0]], table="NAM")


def test_native_roundtrip_preserves_blocks_metadata_scenarios_and_derived_roles(tmp_path):
    source, output = tmp_path / "FIGARO_NAM_IT_2023.xlsx", tmp_path / "native.xlsx"
    write_figaro_source(source)
    db = mario.parse_figaro(source, table="NAM")
    db.clone_scenario("baseline", "policy")
    use = dense(db.get_block("U"))
    use.iloc[0, 0] += 2
    db.set_block("U", use, "policy")
    db.resolve("T", scenario="policy")
    db.resolve_many(["Xa", "u", "cf_d"])
    db.resolve_many(["Xa", "u", "cf_d"], scenario="policy")
    db.to_excel(output)
    restored = mario.parse_from_excel(output, table="NAM", mode="flows")
    assert restored.scenarios == ["baseline", "policy"]
    assert restored.nam_schema.to_dict() == db.nam_schema.to_dict()
    assert restored.info == db.info
    for scenario in db.scenarios:
        assert "T" not in restored.list_matrices(scenario)
        for name in db.list_matrices(scenario):
            if NAM_BLOCKS[name]["role"] == "flow":
                pdt.assert_frame_equal(dense(restored.get_block(name, scenario)), dense(db.get_block(name, scenario)), check_dtype=False)
            else:
                pdt.assert_frame_equal(dense(restored.resolve(name, scenario=scenario)), dense(db.resolve(name, scenario=scenario)), check_dtype=False)
        pdt.assert_frame_equal(dense(restored.resolve("T", scenario=scenario)), dense(db.resolve("T", scenario=scenario)), check_dtype=False)
    copied = restored.copy()
    assert copied.nam_schema.to_dict() == restored.nam_schema.to_dict()
    with pytest.raises(NotImplementable, match="flows only"):
        db.to_excel(output, coefficients=True)


@pytest.mark.parametrize("mode", ["use_imports", "supply_imports"])
def test_template_preserves_unknowns_and_explicit_zeros(mode, tmp_path):
    path = tmp_path / "template.xlsx"
    mario.write_parse_template(path, table="NAM", sets={"IT": {"commodity": ["Good"], "activity": ["Industry"],
                                                                  "external_trade": ["ROW"]}},
                               units={"currency": "EUR", "scale": 1e6, "period": 2023},
                               representation=mode, block_states={"S": "declared_zero"})
    state = read_nam_workbook(path)
    assert state.nam_schema.representation == mode
    assert state.has_matrix("S")
    assert not state.has_matrix("U")
    assert state._resolve_block_record("U").metadata["presence"] == "unavailable"
    assert "MU" in state.nam_schema.required_flows() if mode == "use_imports" else "MS" in state.nam_schema.required_flows()
    with pytest.raises(DataMissing, match="U"):
        state.compute("T")
    wb = load_workbook(path)
    registry = wb["matrix_layouts"]
    row = next(r for r in registry.iter_rows(min_row=2) if r[1].value == "U")
    top = row[4].value
    wb["flows"].cell(top + 3, 4, 12)
    wb.save(path)
    with pytest.raises(WrongInput, match="unavailable block has values"):
        read_nam_workbook(path)
    row[2].value = "observed"
    wb.save(path)
    assert read_nam_workbook(path).get_block("U").iloc[0, 0] == 12


def test_source_malformed_or_unowned_observations_are_not_dropped(tmp_path):
    source = tmp_path / "FIGARO_NAM_IT_2023.xlsx"
    write_figaro_source(source, change=lambda codes, x: x.__setitem__((codes.index("F41"), 0), 1))
    with pytest.raises(WrongInput, match="lack a NAM flow"):
        mario.parse_figaro(source, table="NAM")
    write_figaro_source(source, change=lambda codes, x: x.__setitem__((codes.index("P7"), codes.index("A01")), 11))
    with pytest.raises(WrongInput, match="unexplained P7/detail"):
        mario.parse_figaro(source, table="NAM")
    write_figaro_source(source, change=lambda codes, x: x.__setitem__((0, 0), np.nan))
    with pytest.raises(WrongInput, match="missing/nonfinite"):
        mario.parse_figaro(source, table="NAM")


def test_no_local_path_is_required_for_old_api_modes_but_is_required_for_nam():
    with pytest.raises(WrongInput, match="local workbook"):
        mario.parse_figaro(table="NAM")


def test_native_v2_item_metadata_is_editable_and_v1_is_rejected(tmp_path):
    path = tmp_path / 'template.xlsx'
    mario.write_parse_template(path, table='NAM',
        sets={'IT': {'nf': ['Machinery'], 'o': ['Pool'], 'tr': ['Imports']}},
        units={'currency': 'EUR', 'scale': 1e6, 'period': 2023},
        item_metadata={'IT': {'nf': {'Machinery': {'kind': 'fixed_asset', 'description': 'Equipment', 'source_code': 'N1'}}}})
    state = read_nam_workbook(path)
    assert state.nam_schema.item_metadata['IT']['nonfinancial_asset']['Machinery']['source_code'] == 'N1'
    wb = load_workbook(path)
    row = next(r for r in wb['accounts'].iter_rows(min_row=2) if r[2].value == 'Machinery')
    row[4].value = 'Updated equipment description'
    wb.save(path)
    assert read_nam_workbook(path).nam_schema.item_metadata['IT']['nonfinancial_asset']['Machinery']['description'] == 'Updated equipment description'
    row = next(r for r in wb['metadata'].iter_rows(min_row=2) if r[0].value == 'version')
    row[1].value = 1
    wb.save(path)
    with pytest.raises(WrongInput, match='Version 1 workbooks are not silently reclassified'):
        read_nam_workbook(path)


def write_bulk_source(path, mutation=None):
    codes, main = write_figaro_source(path.with_suffix('.xlsx'))
    aliases = {'CPA_C10': 'CPA_C10-12', 'C10': 'C10-C12'}
    records = []
    for region, multiplier in [('IT', 1), ('FR', 2)]:
        for i, j in zip(*np.nonzero(main)):
            records.append([2023, region, aliases.get(codes[i], codes[i]), region,
                            aliases.get(codes[j], codes[j]), main[i, j] * multiplier])
        for origin, product, destination, amount in [
            ('DE', 'CPA_A01', 'A01', 4), ('DE', 'CPA_A01', 'P51G', 2),
            ('NL', 'CPA_C10-12', 'A01', 6), ('NL', 'CPA_C10-12', 'P3_S14', 5),
        ]:
            records.append([2023, origin, product, region, destination, amount * multiplier])
    # Outgoing bilateral detail must not be mistaken for the IT national ledger
    # or double-counted on top of its observed S2 export totals.
    records.append([2023, 'IT', 'CPA_A01', 'DE', 'A01', 999])
    if mutation:
        mutation(records)
    with path.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.writer(stream)
        writer.writerow(['base', 'ctr', 'Set_i', 'm', 'Set_j', 'value'])
        writer.writerows(records)


def test_bulk_csv_selects_one_national_ledger_and_its_foreign_imports(tmp_path, monkeypatch):
    path = tmp_path / 'FIGARO_NAM_2023.csv'
    write_bulk_source(path)
    from mario.parsers import figaro_nam
    read_csv = figaro_nam.pd.read_csv
    def small_chunks(*args, **kwargs):
        assert kwargs['chunksize'] > 0
        kwargs['chunksize'] = 5
        return read_csv(*args, **kwargs)
    monkeypatch.setattr(figaro_nam.pd, 'read_csv', small_chunks)
    db = mario.parse_figaro(path, table='NAM', countries='IT')
    assert db.get_index('r') == ['IT']
    assert db.get_index('a')[-1] == 'U'
    assert db.get_index('c')[-1] == 'CPA_U'
    assert len(db.get_index('a')) == len(db.get_index('c')) == 64
    assert len(db.get_index('f')) == 6
    assert dense(db.U).loc[('IT', 'Commodity', 'CPA_A01'), ('IT', 'Activity', 'A01')] == 40
    assert dense(db.MU).loc[('IT', 'DE', 'CPA_A01', 'External trade account', 'P7'), ('IT', 'Activity', 'A01')] == 4
    assert dense(db.Yc_exp).to_numpy().sum() == 10
    assert dense(db.TP_trade).to_numpy().sum() == 2
    assert db.info['figaro_nam']['source_format'] == 'bulk_csv'
    assert db.info['figaro_nam']['source_reconciled']
    view = dense(db.T)
    np.testing.assert_allclose(view.sum(axis=0), view.sum(axis=1))
    assert view.loc[('IT', 'Commodity', 'CPA_U')].sum() == 0
    foreign = mario.parse_figaro(path, table='NAM', countries=['FR'], model='ModelState')
    assert foreign.get_block('U').sparse.to_dense().iloc[0, 0] == 80
    assert foreign.nam_schema.period == 2023
    write_bulk_source(path, lambda rows: rows.append(rows[0].copy()))
    with pytest.raises(WrongInput, match='duplicate'):
        mario.parse_figaro(path, table='NAM', countries='IT')


@pytest.mark.parametrize('countries', [None, [], ['IT', 'FR'], ['IT', 'IT'], '', 1])
def test_bulk_csv_requires_explicit_single_country(tmp_path, countries):
    path = tmp_path / 'FIGARO_NAM_2023.csv'
    write_bulk_source(path)
    with pytest.raises(WrongInput, match='exactly one country'):
        mario.parse_figaro(path, table='NAM', countries=countries)


@pytest.mark.parametrize(('mutation', 'message'), [
    (lambda r: r.append(r[0].copy()), 'duplicate'),
    (lambda r: r[0].__setitem__(5, ''), 'nonnumeric'),
    (lambda r: r[0].__setitem__(5, 'inf'), 'nonfinite'),
    (lambda r: r[0].__setitem__(2, 'Unknown'), 'unrecognized'),
    (lambda r: r[0].__setitem__(0, 2022), 'reference year'),
    (lambda r: r.append([2023, 'IT', 'S1', 'IT', 'S14', 1]), 'nonzero S1 control'),
    (lambda r: r.append([2023, 'IT', 'F41', 'IT', 'CPA_A01', 1]), 'lack a NAM flow'),
    (lambda r: r.append([2023, 'DE', 'S14', 'IT', 'A01', 1]), 'foreign bulk rows'),
    (lambda r: r.append([2023, 'DE', 'CPA_A01', 'IT', 'A01', 1]), 'duplicate'),
    (lambda r: r.append([2023, 'ES', 'CPA_A01', 'IT', 'A01', 1]), 'unexplained P7/detail'),
])
def test_bulk_csv_rejects_ambiguous_or_unreconciled_observations(tmp_path, mutation, message):
    path = tmp_path / 'FIGARO_NAM_2023.csv'
    write_bulk_source(path, mutation)
    with pytest.raises(WrongInput, match=message):
        mario.parse_figaro(path, table='NAM', countries='IT')


def test_bulk_csv_selection_year_and_native_roundtrip(tmp_path):
    path, output = tmp_path / 'FIGARO_NAM_2023.csv', tmp_path / 'bulk_native.xlsx'
    write_bulk_source(path)
    with pytest.raises(WrongInput, match='reference year'):
        mario.parse_figaro(path, table='NAM', countries='IT', year=2022)
    with pytest.raises(WrongInput, match='not present'):
        mario.parse_figaro(path, table='NAM', countries='XX')
    with pytest.raises(WrongInput, match='one FIGARO-NAM bulk CSV path'):
        mario.parse_figaro([path, path.with_suffix('.xlsx')], table='NAM', countries='IT')
    db = mario.parse_figaro(path, table='NAM', countries='IT')
    db.to_excel(output)
    restored = mario.parse_from_excel(output, table='NAM', mode='flows')
    assert restored.nam_schema.to_dict() == db.nam_schema.to_dict()
    assert restored.info == db.info
    pdt.assert_frame_equal(dense(restored.T), dense(db.T), check_dtype=False)
    for name in db.list_matrices():
        if name != 'T':
            pdt.assert_frame_equal(dense(restored.get_block(name)), dense(db.get_block(name)), check_dtype=False)
