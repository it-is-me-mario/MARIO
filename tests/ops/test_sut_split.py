"""Tests for the SUT sector split of ``Database.add_sectors``."""

import importlib.util
import warnings

import numpy as np
import pandas as pd
import pytest

from mario.log_exc.exceptions import WrongInput
from mario.ops.add_sector_split_sut import (
    build_sut_split_priors,
    build_sut_split_structure,
    prepare_sut_split_support,
)
from mario.ops.add_sector_workbook import read_add_sector_workbook, validate_sut_add_mode

from .sut_split_helpers import ADD_SECTOR_PATH, VALUE_ADDED, build_example_database, read_expected


@pytest.fixture()
def guide_db():
    db = build_example_database()
    db.read_add_sectors_excel(ADD_SECTOR_PATH, read_inventories=True, split=True)
    return db


# ---------------------------------------------------------------------------
# Workbook reading and add/split mode
# ---------------------------------------------------------------------------
def test_guide_workbook_is_read_as_sut_split(guide_db):
    assert guide_db.to_split_activities == ["a2b"]
    assert guide_db.to_split_commodities == ["c2b"]
    assert guide_db.parented_commodities == ["c2b"]
    assert set(guide_db.split_info) == {
        "Total outputs",
        "Supply coef",
        "Final dem share",
        "Exclusions use",
        "Exclusions supply",
        "Tolerances",
    }


def test_sut_workbook_without_split_columns_is_all_add():
    master = pd.DataFrame({"Activity": ["x"], "Add or Split": [None]})
    assert validate_sut_add_mode(master) == "add"


def test_sut_workbook_cannot_mix_add_and_split():
    master = pd.DataFrame({"Activity": ["x", "y", "z"], "Add or Split": ["Split", "Add", "Split"]})
    with pytest.raises(WrongInput, match="cannot mix"):
        validate_sut_add_mode(master)


def test_sut_workbook_rejects_unknown_mode():
    master = pd.DataFrame({"Activity": ["x"], "Add or Split": ["Merge"]})
    with pytest.raises(WrongInput, match="accepts only"):
        validate_sut_add_mode(master)


def test_mixed_mode_workbook_fails_at_read_time(tmp_path):
    frame = pd.read_excel(ADD_SECTOR_PATH, sheet_name=None)
    frame["Master"].loc[0, "Add or Split"] = "Add"
    path = tmp_path / "mixed.xlsx"
    with pd.ExcelWriter(path) as writer:
        for name, sheet in frame.items():
            sheet.to_excel(writer, sheet_name=name, index=False)
    with pytest.raises(WrongInput, match="cannot mix"):
        read_add_sector_workbook(path, table="SUT", parse_split=True)


def test_optional_split_sheets_can_be_omitted(tmp_path):
    frame = pd.read_excel(ADD_SECTOR_PATH, sheet_name=None)
    path = tmp_path / "minimal.xlsx"
    with pd.ExcelWriter(path) as writer:
        for name, sheet in frame.items():
            if name in {"Exclusions use", "Exclusions supply", "Tolerances"}:
                continue
            sheet.to_excel(writer, sheet_name=name, index=False)
    workbook = read_add_sector_workbook(path, table="SUT", parse_split=True)
    assert workbook.split_info["Exclusions use"].empty
    assert workbook.split_info["Exclusions supply"].empty
    assert workbook.split_info["Tolerances"].empty


def test_required_split_sheet_is_enforced(tmp_path):
    frame = pd.read_excel(ADD_SECTOR_PATH, sheet_name=None)
    path = tmp_path / "missing.xlsx"
    with pd.ExcelWriter(path) as writer:
        for name, sheet in frame.items():
            if name == "Supply coef":
                continue
            sheet.to_excel(writer, sheet_name=name, index=False)
    with pytest.raises(Exception, match="Supply coef"):
        read_add_sector_workbook(path, table="SUT", parse_split=True)


# ---------------------------------------------------------------------------
# Validation of the split-support sheets
# ---------------------------------------------------------------------------
def test_structure_of_the_guide(guide_db):
    structure = build_sut_split_structure(guide_db)
    assert structure.activity_parent == {"a2b": "a2"}
    assert structure.commodity_parent == {"c2b": "c2"}
    assert structure.stable_activities == ["a1"]
    assert structure.parent_commodities == ["c2"]
    assert structure.all_activities == ["a1", "a2", "a2b"]


def test_missing_total_output_is_rejected(guide_db):
    guide_db.split_info["Total outputs"] = guide_db.split_info["Total outputs"].iloc[1:]
    with pytest.raises(WrongInput, match="missing total output"):
        prepare_sut_split_support(guide_db)


def test_new_commodity_in_supply_coef_is_rejected(guide_db):
    sheet = guide_db.split_info["Supply coef"].copy()
    sheet.loc[0, "Commodity_to"] = "c2b"
    guide_db.split_info["Supply coef"] = sheet
    with pytest.raises(WrongInput, match="commodity inventory"):
        prepare_sut_split_support(guide_db)


def test_final_demand_share_outside_unit_interval_is_rejected(guide_db):
    sheet = guide_db.split_info["Final dem share"].copy()
    sheet.loc[0, "Share"] = 1.5
    guide_db.split_info["Final dem share"] = sheet
    with pytest.raises(WrongInput, match="between 0 and 1"):
        prepare_sut_split_support(guide_db)


def test_child_output_larger_than_parent_is_rejected(guide_db):
    sheet = guide_db.split_info["Total outputs"].copy()
    sheet.loc[0, "Quantity"] = 1000.0
    guide_db.split_info["Total outputs"] = sheet
    structure = prepare_sut_split_support(guide_db)
    with pytest.raises(WrongInput, match="larger than parent"):
        build_sut_split_priors(guide_db, structure)


# ---------------------------------------------------------------------------
# Priors: golden test against the guide
# ---------------------------------------------------------------------------
@pytest.fixture()
def guide_priors(guide_db):
    structure = prepare_sut_split_support(guide_db)
    return build_sut_split_priors(guide_db, structure)


def _assert_block(actual: pd.DataFrame, expected: dict, *, label: str):
    for (row, column), value in expected.items():
        got = float(actual.loc[row, column])
        assert got == pytest.approx(value, abs=1e-9), f"{label}[{row}, {column}]: {got} != {value}"


def test_priors_reproduce_the_reference_use(guide_priors):
    expected = read_expected()["Z"]
    # The use of a split commodity by an existing activity is free (solved by CVXLab): compare the
    # rows of unsplit commodities and the columns of the new activity, which the inventories fix.
    rows = [i for i in expected.index if i[1] == "Commodity"]
    cols = [c for c in expected.columns if c[1] == "Activity"]
    for row in rows:
        for col in cols:
            if row[2] != "c1" and col[2] != "a2b":
                continue
            got = float(guide_priors.U0.loc[(row[0], row[2]), (col[0], col[2])])
            assert got == pytest.approx(float(expected.loc[row, col]), abs=1e-9), (row, col)


def test_priors_reproduce_the_reference_final_demand(guide_priors):
    expected = read_expected()["Y"]
    for row in (i for i in expected.index if i[1] == "Commodity"):
        for col in expected.columns:
            got = float(guide_priors.Y0.loc[(row[0], row[2]), (col[0], col[2])])
            assert got == pytest.approx(float(expected.loc[row, col]), abs=1e-9), (row, col)


def test_priors_reproduce_the_reference_outputs(guide_priors):
    assert guide_priors.Xnew[("DE", "a2b")] == pytest.approx(10.8)
    assert guide_priors.Xnew[("IT", "a2b")] == pytest.approx(8.3)
    assert guide_priors.Xnew[("DE", "a2")] == 0.0
    assert guide_priors.Xc[("DE", "c2b")] == pytest.approx(15.0)
    assert guide_priors.Xc[("IT", "c2b")] == pytest.approx(10.0)


def test_priors_zero_mask_follows_exclusions_and_original_zeros(guide_priors):
    zero = guide_priors.zero_M
    # Exclusions supply
    assert bool(zero.loc[("DE", "a1"), ("DE", "c2b")])
    assert bool(zero.loc[("IT", "a2"), ("DE", "c2b")])
    assert bool(zero.loc[("DE", "a2b"), ("DE", "c1")])
    # inherited from the original table (a1 never makes c2 in DE -> DE)
    assert bool(zero.loc[("DE", "a1"), ("DE", "c2")])
    # a cell that is free
    assert not bool(zero.loc[("DE", "a2b"), ("DE", "c2b")])


def test_priors_keep_parent_plus_children_equal_to_original(guide_priors):
    original = guide_priors.original
    U0, M0, Y0, V0 = guide_priors.U0, guide_priors.M0, guide_priors.Y0, guide_priors.V0
    for region in ("DE", "IT"):
        for r2 in ("DE", "IT"):
            # parent activity column of a stable commodity + child column == original
            got = U0.loc[(region, "c1"), (r2, "a2")] + U0.loc[(region, "c1"), (r2, "a2b")]
            assert got == pytest.approx(original.U.loc[(region, "c1"), (r2, "a2")])
            # Make parent cell + children == original
            block = M0.loc[[(region, "a2"), (region, "a2b")], [(r2, "c2"), (r2, "c2b")]].to_numpy().sum()
            assert block == pytest.approx(original.M.loc[(region, "a2"), (r2, "c2")])
            # Final demand
            assert (Y0.loc[(region, "c2")] + Y0.loc[(region, "c2b")]).to_numpy() == pytest.approx(
                original.Y.loc[(region, "c2")].to_numpy()
            )
    assert (V0.loc["Value added", ("DE", "a2")] + V0.loc["Value added", ("DE", "a2b")]) == pytest.approx(12.7)


def test_missing_final_demand_share_warns_and_uses_zero(guide_db):
    guide_db.split_info["Final dem share"] = guide_db.split_info["Final dem share"].iloc[:3]
    structure = prepare_sut_split_support(guide_db)
    with pytest.warns(UserWarning, match="no share"):
        priors = build_sut_split_priors(guide_db, structure)
    assert priors.Y0.loc[("IT", "c2b"), ("IT", "Final demand")] == 0.0
    assert priors.Y0.loc[("IT", "c2"), ("IT", "Final demand")] == pytest.approx(14.1)


def test_priors_clip_negative_parent_with_a_warning(guide_db):
    sheet = guide_db.split_info["Total outputs"].copy()
    # child as large as its parent in use terms makes some parent cells negative
    inventory = guide_db.add_sectors_workbook.inventories_by_sheet["a2b_DE"]
    inventory.loc[inventory["DB Item"] == "c1", "Quantity"] = np.where(
        inventory.loc[inventory["DB Item"] == "c1", "DB Region"] == "DE", 5.0, 0.0
    )
    structure = prepare_sut_split_support(guide_db)
    with pytest.warns(UserWarning, match="clipped"):
        priors = build_sut_split_priors(guide_db, structure)
    assert (priors.U0.to_numpy() >= 0).all()


# ---------------------------------------------------------------------------
# Database.add_sectors integration
# ---------------------------------------------------------------------------
def _cvxlab_available():
    return importlib.util.find_spec("cvxlab") is not None


needs_cvxlab = pytest.mark.skipif(not _cvxlab_available(), reason="CVXLab is not installed")


def test_split_workbook_without_split_flag_is_rejected(tmp_path):
    db = build_example_database()
    with pytest.raises(WrongInput, match="split=True"):
        db.add_sectors(io=ADD_SECTOR_PATH, inplace=False)


def test_add_workbook_with_split_flag_is_rejected(tmp_path):
    frame = pd.read_excel(ADD_SECTOR_PATH, sheet_name=None)
    frame["Master"]["Add or Split"] = "Add"
    path = tmp_path / "all_add.xlsx"
    with pd.ExcelWriter(path) as writer:
        for name, sheet in frame.items():
            sheet.to_excel(writer, sheet_name=name, index=False)
    db = build_example_database()
    with pytest.raises(WrongInput, match="No sectors are marked as 'Split'"):
        db.add_sectors(io=path, split=True, cvxlab_path=tmp_path, inplace=False)


def test_split_requires_cvxlab_path():
    db = build_example_database()
    with pytest.raises(WrongInput, match="cvxlab_path"):
        db.add_sectors(io=ADD_SECTOR_PATH, split=True, inplace=False)


@needs_cvxlab
def test_only_input_data_gen_writes_cvxlab_inputs(tmp_path):
    db = build_example_database()
    new = db.add_sectors(
        io=ADD_SECTOR_PATH, split=True, cvxlab_path=tmp_path, only_input_data_gen=True, inplace=False
    )
    model_dir = tmp_path / "mario_add_sector_split_sut"
    assert (model_dir / "sets.xlsx").exists()
    tables = pd.read_excel(model_dir / "input_data" / "input_data.xlsx", sheet_name=None)
    assert {"Uold", "Mold", "Yold", "Vold", "U0", "M0", "Y0", "V0", "Xnew", "Zero_U", "Zero_M"} <= set(tables)
    identity = tables["I_act_p_pn"].set_index(["activity_to_Name", "activity_from_Name"])["values"]
    assert identity.to_dict() == {("a2", "a2"): 1, ("a2b", "a2"): 1}
    sets = pd.read_excel(model_dir / "sets.xlsx", sheet_name="_set_COMMODITY_FROM")
    assert dict(zip(sets["commodity_from_Name"], sets["commodity_from_category"])) == {
        "c1": "stable",
        "c2": "parent",
        "c2b": "new",
    }
    # without a solve, the prior scenario is promoted to baseline
    assert list(new.scenarios) == ["baseline"]
    assert "a2b" in new.get_index("Activity") and "c2b" in new.get_index("Commodity")


@pytest.fixture(scope="module")
def solved_split(tmp_path_factory):
    pytest.importorskip("cvxlab")
    db = build_example_database()
    folder = tmp_path_factory.mktemp("cvxlab_sut")
    new = db.add_sectors(io=ADD_SECTOR_PATH, split=True, cvxlab_path=folder, inplace=False)
    return db, new


@needs_cvxlab
def test_split_adds_the_new_items(solved_split):
    original, new = solved_split
    assert new.get_index("Activity") == ["a1", "a2", "a2b"]
    assert new.get_index("Commodity") == ["c1", "c2", "c2b"]
    assert new.units["Activity"].loc["a2b", "unit"] == "M EUR"
    assert list(new.scenarios) == ["baseline"]


@needs_cvxlab
def test_split_keeps_parent_plus_children_equal_to_original(solved_split):
    original, new = solved_split
    U_old, S_old = original.U, original.S
    U_new, S_new = new.U, new.S
    for region in ("DE", "IT"):
        for r2 in ("DE", "IT"):
            for commodity in ("c1",):
                total = (
                    U_new.loc[(region, "Commodity", commodity), (r2, "Activity", "a2")]
                    + U_new.loc[(region, "Commodity", commodity), (r2, "Activity", "a2b")]
                )
                assert total == pytest.approx(U_old.loc[(region, "Commodity", commodity), (r2, "Activity", "a2")], rel=1e-4)
            block = S_new.loc[
                [(region, "Activity", "a2"), (region, "Activity", "a2b")],
                [(r2, "Commodity", "c2"), (r2, "Commodity", "c2b")],
            ].to_numpy().sum()
            assert block == pytest.approx(S_old.loc[(region, "Activity", "a2"), (r2, "Commodity", "c2")], rel=1e-4)
    # stable cells are untouched
    assert S_new.loc[("DE", "Activity", "a1"), ("DE", "Commodity", "c1")] == pytest.approx(14.4)


@needs_cvxlab
def test_split_is_balanced_and_respects_known_outputs(solved_split):
    _, new = solved_split
    X = new.X["production"]
    assert X[("DE", "Activity", "a2b")] == pytest.approx(10.8, rel=1e-4)
    assert X[("IT", "Activity", "a2b")] == pytest.approx(8.3, rel=1e-4)
    # outputs are consistent with the aggregated ones
    assert X[("DE", "Activity", "a2")] + X[("DE", "Activity", "a2b")] == pytest.approx(35.7, rel=1e-4)
    assert X[("DE", "Commodity", "c2")] + X[("DE", "Commodity", "c2b")] == pytest.approx(37.0, rel=1e-4)
    # activity balance: purchases + value added = output
    U, V = new.U, new.V
    for region in ("DE", "IT"):
        for activity in ("a1", "a2", "a2b"):
            column = (region, "Activity", activity)
            assert U[column].sum() + V[column].sum() == pytest.approx(X[column], rel=1e-4)


@needs_cvxlab
def test_split_stays_close_to_the_reference_disaggregation(solved_split):
    """The reference table is the ground truth the inventories were derived from.

    The optimization only sees the priors, so it recovers the reference approximately: the Make
    table within 0.5, final demand exactly (it is fixed by the shares), and the value added of the
    new activities as the output not paid to inputs.
    """

    _, new = solved_split
    reference = read_expected()
    Z = reference["Z"]
    make = Z.loc[[i for i in Z.index if i[1] == "Activity"], [c for c in Z.columns if c[1] == "Commodity"]]
    assert np.abs(new.S.to_numpy() - make.to_numpy()).max() < 0.5
    final = reference["Y"].loc[[i for i in reference["Y"].index if i[1] == "Commodity"]]
    assert np.abs(new.Y.loc[final.index, final.columns].to_numpy() - final.to_numpy()).max() < 1e-4
    va = new.V.loc[VALUE_ADDED]
    for region, expected in (("DE", 1.1), ("IT", 1.3)):
        assert va[(region, "Activity", "a2b")] == pytest.approx(expected, abs=1.0)
    assert va[("DE", "Activity", "a2")] + va[("DE", "Activity", "a2b")] == pytest.approx(12.7, rel=1e-4)


@needs_cvxlab
def test_split_respects_exclusions(solved_split):
    _, new = solved_split
    S = new.S
    assert S.loc[("DE", "Activity", "a1"), ("DE", "Commodity", "c2")] == pytest.approx(0, abs=1e-6)
    assert S.loc[("DE", "Activity", "a2b"), ("DE", "Commodity", "c1")] == pytest.approx(0, abs=1e-6)
    assert S.loc[("IT", "Activity", "a2"), ("DE", "Commodity", "c2b")] == pytest.approx(0, abs=1e-6)
    assert (new.U.to_numpy() >= -1e-9).all()


@needs_cvxlab
def test_split_satellite_accounts_follow_parent_intensity(solved_split):
    original, new = solved_split
    E_old, E_new = original.E, new.E
    # the new activity inherits the parent intensity and the total is preserved
    total = E_new.loc["CO2", ("DE", "Activity", "a2")] + E_new.loc["CO2", ("DE", "Activity", "a2b")]
    assert total == pytest.approx(E_old.loc["CO2", ("DE", "Activity", "a2")])
    ratio = E_new.loc["CO2", ("DE", "Activity", "a2b")] / new.X["production"][("DE", "Activity", "a2b")]
    assert ratio == pytest.approx(E_old.loc["CO2", ("DE", "Activity", "a2")] / 35.7)


@needs_cvxlab
def test_split_keep_all_steps_and_parent_rename(tmp_path):
    db = build_example_database()
    new = db.add_sectors(
        io=ADD_SECTOR_PATH,
        split=True,
        cvxlab_path=tmp_path,
        keep_all_split_steps=True,
        parent_names={"Activity": {"a2": "a2_rest"}, "Commodity": {"c2": "c2_rest"}},
        inplace=False,
    )
    assert set(new.scenarios) == {"baseline", "original", "split_baseline", "split_cvxlab"}
    assert new.get_index("Activity") == ["a1", "a2_rest", "a2b"]
    assert new.get_index("Commodity") == ["c1", "c2_rest", "c2b"]
    assert "a2_rest" in new.units["Activity"].index
    assert ("DE", "Activity", "a2_rest") in new.X.index


# ---------------------------------------------------------------------------
# Templates and satellite inventories
# ---------------------------------------------------------------------------
def test_sut_template_gets_split_sheets_when_master_is_marked_split(tmp_path):
    from openpyxl import load_workbook

    db = build_example_database()
    path = tmp_path / "template.xlsx"
    db.get_add_sectors_excel(path=path, items=["x"], regions=["DE"], item="Activity")
    master = pd.read_excel(path, sheet_name="Master")
    assert {"Parent Commodity", "Add or Split"} <= set(master.columns)
    assert master["Add or Split"].tolist() == ["Add"]

    master["Add or Split"] = "Split"
    with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        master.to_excel(writer, sheet_name="Master", index=False)
    db.read_add_sectors_excel(path)
    db.get_inventory_sheets(path)
    sheets = set(load_workbook(path).sheetnames)
    assert {"Total outputs", "Supply coef", "Final dem share", "Exclusions use", "Exclusions supply", "Tolerances"} <= sheets
    assert "Trades" not in sheets


def test_satellite_inventory_rows_override_parent_intensity(guide_db):
    inventories = guide_db.add_sectors_workbook.inventories_by_sheet
    extra = pd.DataFrame(
        [{"Quantity": 10.0, "Unit": "kg", "Input": None, "Item type": "Satellite account", "DB Item": "CO2",
          "DB Region": None, "Change type": "Update"}]
    )
    inventories["a2b_DE"] = pd.concat([inventories["a2b_DE"], extra], ignore_index=True)
    structure = prepare_sut_split_support(guide_db)
    priors = build_sut_split_priors(guide_db, structure)
    from mario.ops.add_sector_split_sut import assemble_sut_split_flows

    flows = assemble_sut_split_flows(
        structure, priors.original, priors,
        U=priors.U0, M=priors.M0, Y=priors.Y0, V=priors.V0, Xa=priors.Xa, Xc=priors.Xc,
    )
    E = flows["E"]
    # 10 kg per unit of the 10.8 output of a2b in DE; IT keeps the parent intensity
    assert E.loc["CO2", ("DE", "Activity", "a2b")] == pytest.approx(108.0)
    assert E.loc["CO2", ("IT", "Activity", "a2b")] == pytest.approx(1.0 / 35 * 8.3)


# ---------------------------------------------------------------------------
# Unknown items are rejected
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "sheet, column, value",
    [
        ("Supply coef", "Commodity_to", "nope"),
        ("Supply coef", "Activity_from", "nope"),
        ("Exclusions use", "Commodity_from", "nope"),
        ("Exclusions use", "Activity_to", "nope"),
        ("Exclusions supply", "Activity_from", "nope"),
        ("Exclusions supply", "Commodity_to", "nope"),
        ("Supply coef", "Region_to", "XX"),
    ],
)
def test_unknown_items_in_split_sheets_are_rejected(guide_db, sheet, column, value):
    frame = guide_db.split_info[sheet].copy()
    frame.loc[0, column] = value
    guide_db.split_info[sheet] = frame
    with pytest.raises(WrongInput, match="unknown"):
        prepare_sut_split_support(guide_db)
