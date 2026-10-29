"""Saved-point restoration must never solve or silently accept changed inputs."""
import json
import shutil
import sqlite3
from types import SimpleNamespace

import cvxpy as cp
import numpy as np
import pandas as pd
import pytest

from mario.log_exc.exceptions import WrongInput
from mario.ops.split_solution_replay import replay_split_solution, restore_saved_solution


def fixture(tmp_path):
    source, target = tmp_path/'source', tmp_path/'target'
    source.mkdir()
    (source/'input_data').mkdir()
    for name in ['model_settings.xlsx', 'mapping.xlsx', 'user_defined_operators.py', 'user_defined_constants.py']:
        (source/name).write_bytes(b'fixed model')
    (source/'input_data/prior.csv').write_text('id,values\n1,1.0\n')
    with sqlite3.connect(source/'database.db') as conn:
        pd.DataFrame(dict(id=[1], name=['R'])).to_sql('_set_REGION', conn, index=False)
        pd.DataFrame(dict(id=[1], region=['R'], values=[None])).to_sql('Z_use', conn, index=False)
    (source/'raw_model_solution.json').write_text(json.dumps([dict(status='optimal', solver_seconds=34.)]))
    (source/'returned_solution_index.json').write_text(json.dumps([dict(name='Z_use', array='variable_0')]))
    np.savez_compressed(source/'returned_solution_scaled.npz', variable_0=np.array([[-1e-9]]))
    shutil.copytree(source, target)
    variable = cp.Variable((1, 1), name='Z_use')
    problem = cp.Problem(cp.Minimize(cp.sum(variable)), [variable >= 0])
    frame = pd.DataFrame([dict(problem=problem, status=None)])
    core = SimpleNamespace(numerical_problems={None: frame}, problem_status=None)
    core.fetch_problem_status = lambda: setattr(core, 'problem_status', {'saved': frame.at[0, 'status']})
    return source, target, SimpleNamespace(core=SimpleNamespace(problem=core)), variable


def test_restoration_reuses_the_raw_point_once_without_a_solver(tmp_path, monkeypatch):
    source, target, model, variable = fixture(tmp_path)
    monkeypatch.setattr(cp.Problem, 'solve', lambda *a, **k: pytest.fail('No optimizer may run.'))
    assert not restore_saved_solution(model, target)
    with replay_split_solution(source) as state:
        assert restore_saved_solution(model, target)
        assert state['used']
        assert variable.value[0, 0] == -1e-9
        assert model.core.problem.problem_status == {'saved': 'optimal'}
        with pytest.raises(WrongInput, match='only one split step'):
            restore_saved_solution(model, target)
    record = json.loads((target/'solver_replay.json').read_text())
    assert not record['optimization_rerun'] and record['original_solver_seconds'] == 34.
    assert not restore_saved_solution(model, target)


@pytest.mark.parametrize('change', ['input', 'coordinate', 'status'])
def test_restoration_rejects_incompatible_inputs_or_status(tmp_path, change):
    source, target, model, variable = fixture(tmp_path)
    if change == 'input':
        (target/'input_data/prior.csv').write_text('id,values\n1,2.0\n')
    elif change == 'coordinate':
        with sqlite3.connect(target/'database.db') as conn:
            conn.execute('UPDATE Z_use SET region=?', ['another'])
    else:
        (source/'raw_model_solution.json').write_text(json.dumps([dict(status='optimal_inaccurate')]))
    with pytest.raises(WrongInput):
        with replay_split_solution(source):
            restore_saved_solution(model, target)
    assert variable.value is None


@pytest.mark.parametrize('change_cell', [False, True])
def test_excel_document_metadata_is_ignored_but_model_cells_are_verified(tmp_path, change_cell):
    from datetime import datetime
    from openpyxl import Workbook, load_workbook
    source, target, model, _ = fixture(tmp_path)
    book = Workbook()
    book.active['A1'] = 'fixed expression'
    book.save(source/'model_settings.xlsx')
    book = load_workbook(source/'model_settings.xlsx')
    book.properties.title = 'Regenerated document metadata'
    book.properties.modified = datetime(2026, 10, 23)
    if change_cell:
        book.active['A1'] = 'different expression'
    book.save(target/'model_settings.xlsx')
    if change_cell:
        with pytest.raises(WrongInput, match='model_settings.xlsx'):
            with replay_split_solution(source):
                restore_saved_solution(model, target)
    else:
        with replay_split_solution(source):
            assert restore_saved_solution(model, target)
