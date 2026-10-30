"""Inventory-driven plans, hard scheduling rules and lossless workbook updates."""
import hashlib
import json
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

import mario
from mario.log_exc.exceptions import WrongInput
from mario.ops import add_sector_plan as planning
from mario.ops.add_sector_workbook import read_add_sector_workbook
from tests.ops.test_add_sector_steps import balanced_database, workbook


def source(tmp_path, *, names=("Cars", "Batteries", "Cathode"), recipes=None,
           outputs=None, parents=None, clusters=None, with_plan=True):
    db = balanced_database()
    path = tmp_path / 'inputs.xlsx'
    db.get_add_sectors_excel(path=path, items=list(names), regions=['R1'], intermediate_steps=with_plan)
    master = pd.read_excel(path, sheet_name='Master').astype(object)
    master['Parent Sector'] = [parents.get(s, 'Parent') if parents else 'Parent' for s in master.Sector]
    master['Unit'], master['Add or Split'] = 'M USD', 'Split'
    recipes = recipes if recipes is not None else {
        buyer:[(supplier,q) for supplier,q in rows if supplier in names]
        for buyer,rows in {'Cars':[('Batteries',.3)], 'Batteries':[('Cathode',.2)]}.items()}
    with pd.ExcelWriter(path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        master.to_excel(writer, sheet_name='Master', index=False)
        for record in master.to_dict('records'):
            rows = [[q, 'M USD', item, 'Sector', item, 'R1', 'Update']
                    for item, q in recipes.get(record['Sector'], [])]
            rows.append([.5, 'M USD', 'Value added', 'Factor of production', 'Value added', None, 'Update'])
            pd.DataFrame(rows, columns=['Quantity','Unit','Input','Item type','DB Item','DB Region','Change type']).to_excel(
                writer, sheet_name=record['Inventory sheet'], index=False)
        pd.DataFrame([[s, 'R1', (outputs or {}).get(s, 10.), 'M USD', 'test', None] for s in names],
                     columns=['Sector','Region','Quantity','Unit','Source','Notes']).to_excel(
                         writer, sheet_name='Total outputs', index=False)
        pd.DataFrame({'tol_Name':['delta','eps'], 'values':[0.,0.]}).to_excel(
            writer, sheet_name='Tolerances', index=False)
        if clusters:
            pd.DataFrame({k:pd.Series(v) for k,v in clusters.items()}).to_excel(
                writer, sheet_name='Sectors Clusters', index=False)
    return db, path


def flattened(plan):
    return [s for group in plan.steps.values() for s in group]


def parts(path):
    with ZipFile(path) as archive:
        return {name:archive.read(name) for name in archive.namelist()}


@pytest.mark.parametrize('ordering,expected', [
    ('suppliers_first', ['Cathode','Batteries','Cars']),
    ('buyers_first', ['Cars','Batteries','Cathode']),
])
def test_chain_order_and_exact_preservation_of_other_workbook_parts(tmp_path, ordering, expected, monkeypatch):
    db, path = source(tmp_path)
    before = parts(path)
    before_z = db.Z.copy()
    monkeypatch.setattr(mario.Database, 'add_sectors', lambda *a, **k: pytest.fail('Planning invoked a solve'))
    plan = db.get_add_sectors_intermediate_steps_plan(path, ordering=ordering)
    assert flattened(plan) == expected
    assert plan.summary.sector_count.tolist() == [1,1,1]
    assert plan.diagnostics['cycles'] == [] and plan.diagnostics['unmet_preferences'] == []
    assert plan.path == path.resolve()
    after = parts(path)
    changed = [name for name in before if before[name] != after[name]]
    assert len(changed) == 1 and changed[0].startswith('xl/worksheets/')
    assert before.keys() == after.keys()
    assert read_add_sector_workbook(path, table='IOT').intermediate_steps == plan.steps
    assert db.add_sectors_workbook.intermediate_steps == plan.steps
    pd.testing.assert_frame_equal(db.Z, before_z, check_exact=True)


def test_missing_sheet_is_added_and_formulas_and_unrelated_parts_survive(tmp_path):
    db, path = source(tmp_path, with_plan=False)
    book = load_workbook(path)
    book.create_sheet('User notes')['A1'] = '=SUM(1,2)'
    book['User notes']['B1'] = 'Keep this note'
    book.save(path); book.close()
    before = parts(path)
    plan = db.get_add_sectors_intermediate_steps_plan(path)
    after = parts(path)
    metadata = {'xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml'}
    assert all(after[k] == value for k,value in before.items() if k not in metadata)
    book = load_workbook(path, data_only=False)
    assert book['User notes']['A1'].value == '=SUM(1,2)'
    assert book['Intermediate steps']['A2'].value == 'Cathode'
    book.close()
    assert flattened(plan) == ['Cathode','Batteries','Cars']


def test_weighted_cycle_breaking_uses_production_and_reports_unmet_preferences(tmp_path):
    db, path = source(tmp_path, names=('A','B'), recipes={'A':[('B',.2)], 'B':[('A',.8)]},
                      outputs={'A':100., 'B':1.})
    plan = db.get_add_sectors_intermediate_steps_plan(path)
    assert flattened(plan) == ['B','A']
    assert plan.diagnostics['cycles'] == [['A','B']]
    assert plan.diagnostics['unmet_preferences'][0]['weight'] == pytest.approx(.8)
    assert plan.diagnostics['globally_optimal'] is False
    again = db.get_add_sectors_intermediate_steps_plan(path)
    assert again.steps == plan.steps and again.diagnostics == plan.diagnostics


def test_groups_maximum_and_strict_user_precedence_override_soft_links(tmp_path):
    db, path = source(tmp_path)
    plan = db.get_add_sectors_intermediate_steps_plan(path, max_sectors_per_step=2,
        together=[['Cathode','Batteries']], before=[('Batteries','Cars')])
    assert list(plan.steps.values()) == [['Batteries','Cathode'], ['Cars']]
    plan = db.get_add_sectors_intermediate_steps_plan(path, max_sectors_per_step=3,
        before=[('Cars','Batteries')])
    step = {s:i for i,group in enumerate(plan.steps.values()) for s in group}
    assert step['Cars'] < step['Batteries']
    assert any(r['supplier']=='Batteries' and r['buyer']=='Cars'
               for r in plan.diagnostics['unmet_preferences'])


def test_parent_must_precede_child_even_with_reverse_ordering(tmp_path):
    db, path = source(tmp_path, names=('Cars','Batteries'), parents={'Cars':'Batteries'})
    plan = db.get_add_sectors_intermediate_steps_plan(path, max_sectors_per_step=2, ordering='buyers_first')
    assert list(plan.steps.values()) == [['Batteries'],['Cars']]


@pytest.mark.parametrize('kwargs,message', [
    ({'max_sectors_per_step':0}, 'positive integer'),
    ({'max_sectors_per_step':True}, 'positive integer'),
    ({'max_sectors_per_step':1.5}, 'positive integer'),
    ({'ordering':'magic'}, 'ordering'),
    ({'before':[('Cars','Missing')]}, 'before requires'),
    ({'before':[('Cars','Cars')]}, 'before itself'),
    ({'before':[('Cars','Batteries'),('Batteries','Cars')]}, 'Contradictory'),
    ({'together':[['Cars','Batteries']]}, 'exceeds'),
    ({'together':[['Missing']]}, 'together requires'),
    ({'max_sectors_per_step':2, 'before':[('Cars','Batteries')],
      'together':[['Cars','Batteries']]}, 'conflicts'),
])
def test_invalid_constraints_leave_original_workbook_and_database_untouched(tmp_path, kwargs, message):
    db, path = source(tmp_path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    with pytest.raises(WrongInput, match=message):
        db.get_add_sectors_intermediate_steps_plan(path, **kwargs)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    assert not hasattr(db, 'add_sectors_workbook')


def test_singleton_clusters_identify_links_but_multi_item_clusters_do_not(tmp_path):
    db, path = source(tmp_path, names=('A','B','C'), recipes={'A':[('OnlyB',.3)],'B':[('Broad',.2)]},
        clusters={'OnlyB':['B'], 'Broad':['A','C']})
    plan = db.get_add_sectors_intermediate_steps_plan(path)
    assert {(r['supplier'],r['buyer']) for r in plan.diagnostics['dependencies']} == {('B','A')}
    assert plan.diagnostics['ambiguous_clusters'][0]['members'] == ['A','C']


def test_parent_only_input_does_not_invent_a_battery_link(tmp_path):
    db, path = source(tmp_path, names=('Cars','Batteries'), recipes={'Cars':[('Stable',.3)]})
    plan = db.get_add_sectors_intermediate_steps_plan(path)
    assert plan.diagnostics['dependencies'] == []
    assert {r['sector'] for r in plan.diagnostics['profiles_without_explicit_dependencies']} == {'Cars','Batteries'}


def test_units_are_converted_before_weighting(tmp_path):
    db, path = source(tmp_path, names=('Cars','Batteries'), recipes={'Cars':[('Batteries',.3)]})
    book = load_workbook(path)
    inventory = book['INV_001']
    inventory['A2'], inventory['B2'] = 300000., 'USD'
    output = book['Total outputs']
    output['C2'], output['D2'] = 100000000., 'USD'
    book.save(path); book.close()
    plan = db.get_add_sectors_intermediate_steps_plan(path)
    assert plan.diagnostics['dependencies'][0]['weight'] == pytest.approx(30.)


def test_default_uses_last_loaded_file_and_a_failed_atomic_write_preserves_it(tmp_path, monkeypatch):
    db, path = source(tmp_path)
    db.read_add_sectors_excel(path, read_inventories=True, split=True)
    original = path.read_bytes()
    old_steps = db.add_sectors_workbook.intermediate_steps.copy()
    monkeypatch.setattr(planning.os, 'replace', lambda *a: (_ for _ in ()).throw(OSError('disk failure')))
    with pytest.raises(OSError, match='disk failure'):
        db.get_add_sectors_intermediate_steps_plan()
    assert path.read_bytes() == original
    assert db.add_sectors_workbook.intermediate_steps == old_steps
    assert not list(tmp_path.glob('.*-steps-*.xlsx'))


def test_returned_dictionary_is_detached_from_loaded_executable_plan(tmp_path):
    db, path = source(tmp_path)
    plan = db.get_add_sectors_intermediate_steps_plan(path)
    original = {name:list(group) for name,group in plan.steps.items()}
    plan.steps['Step 1'].clear()
    assert db.add_sectors_workbook.intermediate_steps == original
    assert read_add_sector_workbook(path, table='IOT').intermediate_steps == original


def test_concurrent_edit_is_not_overwritten(tmp_path, monkeypatch):
    db, path = source(tmp_path)
    extract = planning._dependencies
    edited = []
    def concurrent(*args):
        result = extract(*args)
        with path.open('ab') as stream:stream.write(b'External edit')
        edited.append(path.read_bytes())
        return result
    monkeypatch.setattr(planning, '_dependencies', concurrent)
    with pytest.raises(WrongInput, match='changed during planning'):
        db.get_add_sectors_intermediate_steps_plan(path)
    assert path.read_bytes() == edited[0]


def test_checkpoint_inputs_are_immutable(tmp_path):
    db, path = source(tmp_path)
    (tmp_path/'state.json').write_text(json.dumps({'format':'mario-sector-steps-v1','files':{path.name:'hash'}}))
    with pytest.raises(WrongInput, match='outside its committed checkpoint'):
        db.get_add_sectors_intermediate_steps_plan(path)


def test_add_sectors_executes_manually_edited_sheet_without_replanning(tmp_path, monkeypatch):
    db = balanced_database()
    path = workbook(db, tmp_path/'inputs.xlsx')
    db.get_add_sectors_intermediate_steps_plan(path)
    with pd.ExcelWriter(path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        pd.DataFrame({'Manual first':['Mining'], 'Manual second':['Refining']}).to_excel(
            writer, sheet_name='Intermediate steps', index=False)
    original_add = mario.Database.add_sectors
    attempted = []
    def fail_solver(self, *args, **kwargs):
        if not kwargs.get('intermediate_steps'):
            attempted.append(list(self.add_sectors_master.Sector))
            raise WrongInput('CVXLab split optimization did not solve optimally (synthetic)')
        return original_add(self, *args, **kwargs)
    monkeypatch.setattr(mario.Database, 'add_sectors', fail_solver)
    monkeypatch.setattr(mario.Database, 'get_add_sectors_intermediate_steps_plan',
                        lambda *a, **k: pytest.fail('Execution recalculated the plan'))
    db.add_sectors(path, split=True, intermediate_steps=True, on_step_failure='continue',
                   save_intermediates=tmp_path/'checkpoints')
    assert attempted == [['Mining'],['Refining']]


def test_missing_file_and_unloaded_workbook_give_actionable_errors(tmp_path):
    db = balanced_database()
    with pytest.raises(WrongInput, match='first load'):
        db.get_add_sectors_intermediate_steps_plan()
    with pytest.raises(WrongInput, match='existing .xlsx'):
        db.get_add_sectors_intermediate_steps_plan(tmp_path)
