"""Block bridge equivalence, signed cells and sparse numerical constants."""
from types import SimpleNamespace
import sqlite3
import shutil
from zipfile import ZipFile

import cvxpy as cp
import numpy as np
import pandas as pd
from pandas.testing import assert_frame_equal
from scipy import sparse

from mario.ops import cvxlab_bridge as bridge
from mario.ops.cvxlab_models.Split_sectors.user_defined_operators import block_diag
from mario.ops.cvxlab_models.Split_sectors_MRIO_KL.user_defined_constants import my_identity_matrix


def fixture():
    def axis(sectors):
        return pd.MultiIndex.from_product([['B', 'A'], ['Sector'], sectors], names=['Region', 'Level', 'Item'])
    old, new = axis(['Stay', 'Old']), axis(['Stay', 'Old', 'New'])
    demand = pd.MultiIndex.from_product([['B', 'A'], ['Consumption category'], ['H', 'G']], names=old.names)
    z = pd.DataFrame(np.arange(16, dtype=float).reshape(4, 4), index=old, columns=old)
    y = pd.DataFrame(np.arange(16, dtype=float).reshape(4, 4)-3, index=old, columns=demand)
    v = pd.DataFrame(np.arange(16, dtype=float).reshape(4, 4)-5,
                     index=pd.Index(['Tax', 'Capital', 'TTM', 'Labor'], name='Factor of production'), columns=old)
    znew = z.reindex(index=new, columns=new, fill_value=0.)
    hit = new.get_level_values(-1).isin(['Old', 'New'])
    znew.iloc[hit, :] *= .6
    znew.iloc[:, hit] *= .7
    znew.iloc[-1, -1] = 3.
    blocks = {'original': {'Z': z, 'Y': y, 'V': v, 'X': z.sum(axis=1).to_frame('production')},
              'split_baseline': {'Z': znew, 'Y': y.reindex(new, fill_value=0.),
                                 'V': v.reindex(columns=new, fill_value=0.),
                                 'X': znew.sum(axis=1).to_frame('production')}}
    return SimpleNamespace(matrices=blocks, get_block=lambda n, scenario: blocks[scenario][n],
        has_matrix=lambda n, scenario: n in blocks[scenario],
        get_block_as_pandas=lambda n, scenario: blocks[scenario][n].copy()), blocks


def test_filtered_bridge_retains_all_needed_cells_and_immutable_sources():
    db, sources = fixture()
    snapshots = {(s, n): f.copy() for s, b in sources.items() for n, f in b.items()}
    full = bridge._collect_flat_matrices(db, scenarios=list(sources))
    thin = bridge._collect_flat_matrices(db, scenarios=list(sources), affected_sectors={'Old', 'New'})
    for n in full:
        expected = full[n]
        if n == 'Z':
            expected = expected[expected.Sector_from.isin(['Old', 'New']) | expected.Sector_to.isin(['Old', 'New'])]
        elif n == 'Y':
            expected = expected[expected.Sector_from.isin(['Old', 'New'])]
        elif n == 'V':
            expected = expected[expected.Sector_to.isin(['Old', 'New'])]
        cols = list(expected.columns)
        assert_frame_equal(expected.sort_values(cols).reset_index(drop=True),
                           thin[n].sort_values(cols).reset_index(drop=True))
    assert len(thin['Z']) < len(full['Z'])
    for (s, n), snapshot in snapshots.items():
        assert_frame_equal(snapshot, sources[s][n])
        assert thin.source_blocks[s][n] is sources[s][n]


def test_block_reconstruction_matches_legacy_melt_pivot_and_keeps_signed_cells(tmp_path):
    db, sources = fixture()
    full = bridge._collect_flat_matrices(db, scenarios=list(sources))
    thin = bridge._collect_flat_matrices(db, scenarios=list(sources), affected_sectors={'Old', 'New'})
    mapping = {'Region_from': 'region_from_Name', 'Sector_from': 'sector_from_Name',
               'Region_to': 'region_to_Name', 'Sector_to': 'sector_to_Name', 'Value': 'values',
               'Consumption category_to': 'cons_categ_Name', 'Factor of production_from': 'factor_Name'}
    tables = {n: full[n].query("Scenario == 'split_baseline'").drop(columns='Scenario').rename(columns=mapping)
              for n in ('Z', 'Y', 'V')}
    pn = ['Old', 'New']
    sql = {'Z_use': tables['Z'][tables['Z'].sector_to_Name.isin(pn)],
           'Z_supply': tables['Z'][tables['Z'].sector_from_Name.isin(pn) & ~tables['Z'].sector_to_Name.isin(pn)],
           'Y': tables['Y'][tables['Y'].sector_from_Name.isin(pn)],
           'V': tables['V'][tables['V'].sector_to_Name.isin(pn)]}
    with sqlite3.connect(tmp_path / bridge.cl.Defaults.ConfigFiles.SQLITE_DATABASE_FILE) as conn:
        for n, f in sql.items():
            f.reset_index(drop=True).rename_axis('id').reset_index().to_sql(n, conn, index=False)
    asset = bridge.resources.files(bridge.CVXLAB_TEMPLATE_PACKAGE).joinpath('Split_sectors_MRIO_KL/mapping.xlsx')
    shutil.copy2(asset, tmp_path / 'mapping.xlsx')
    with pd.ExcelWriter(tmp_path / bridge.cl.Defaults.ConfigFiles.SETS_FILE) as w:
        pd.DataFrame({'sector_from_Name': ['Stay', 'Old', 'New'],
                      'sector_from_category': ['stable', 'parent', 'new']}).to_excel(w, sheet_name='_set_SECTOR_FROM', index=False)
    snapshots = {n: f.copy() for n, f in sources['original'].items()}
    legacy = bridge._parse_split_results(tmp_path, full)
    thin.clear()  # bridge rows are released before optimisation
    result = bridge._parse_split_results(tmp_path, thin)
    for n in result:
        assert_frame_equal(legacy[n], result[n])
    assert (result['Y'].to_numpy() < 0).any() and (result['V'].to_numpy() < 0).any()
    for n, f in snapshots.items():
        assert_frame_equal(f, sources['original'][n])


def test_sparse_constants_have_identical_algebra_without_dense_expansion():
    a = cp.Constant(np.array([[1., 2.], [0., 3.]]))
    result = block_diag(a, cp.Constant(np.array([[3.]])))
    assert sparse.issparse(result.value)
    np.testing.assert_array_equal(result.value.toarray(), np.kron(np.eye(3), a.value))
    identity = my_identity_matrix([5, 5])
    np.testing.assert_array_equal(identity, np.eye(5))


def test_engine_reuses_numeric_buffers_without_changing_reference_axes(monkeypatch):
    from mario.ops.add_sector_engine import AddSectorEngine, run_add_sector_engine
    index = pd.MultiIndex.from_product([['R'], ['S']], names=['Region', 'Sector'])
    demand = pd.MultiIndex.from_product([['R'], ['H']], names=['Region', 'Consumption category'])
    reference = {'z': pd.DataFrame([[.2]], index=index, columns=index),
                 'Y': pd.DataFrame([[8.]], index=index, columns=demand)}
    snapshots = {n: f.copy() for n, f in reference.items()}
    db = SimpleNamespace(inventories={}, units={}, meta=SimpleNamespace(table='IOT'),
                         table_type='IOT', get_index=lambda name: ['R'] if name == 'Region' else ['S'])
    monkeypatch.setattr(AddSectorEngine, 'to_iot', lambda self, **kwargs: self.matrices)
    result = run_add_sector_engine(db, reference_matrices=reference)
    for n, f in reference.items():
        assert_frame_equal(f, snapshots[n])
        assert np.shares_memory(result[n].to_numpy(), f.to_numpy())
    assert result['Y'].index.nlevels == 3


def test_streamed_workbook_matches_formula_reader_with_bad_dimensions_and_empty_cells(tmp_path):
    from openpyxl import Workbook
    from mario.ops.add_sector_workbook import _read_excel_workbook, _read_excel_workbook_with_formulas
    path = tmp_path / 'values.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.append(['Name', 'Value', 'Optional'])
    sheet.append(['Negative', -3.25])
    sheet.append([None, None, 'Tail'])
    sheet.append([True, 0, None])
    optional = book.create_sheet('Blank optional')
    optional.append(['Name', 'Value', 'Optional'])
    optional.append(['Example', 2])
    book.create_sheet('Empty')
    book.save(path)
    damaged = tmp_path / 'producer_dimensions.xlsx'
    with ZipFile(path) as source, ZipFile(damaged, 'w') as output:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'ref="A1:C4"', b'ref="A1:A1"')
            output.writestr(member, data)
    expected = _read_excel_workbook_with_formulas(path)
    for file in (path, damaged):
        actual = _read_excel_workbook(file)
        assert actual.keys() == expected.keys()
        for name in expected:
            assert_frame_equal(actual[name], expected[name])


def test_streamed_reader_retains_cross_sheet_formula_evaluation(tmp_path):
    from openpyxl import Workbook
    from mario.ops.add_sector_workbook import _read_excel_workbook, _read_excel_workbook_with_formulas
    path = tmp_path / 'formulas.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Source'
    sheet.append(['Value'])
    sheet.append([4.])
    result = book.create_sheet('Result')
    result.append(['Value'])
    result.append(["='Source'!A2*3"])
    result.append(['=SUM(A2,2)'])
    book.save(path)
    actual, expected = _read_excel_workbook(path), _read_excel_workbook_with_formulas(path)
    for name in expected:
        assert_frame_equal(actual[name], expected[name])
    assert actual['Result']['Value'].tolist() == [12., 14.]
