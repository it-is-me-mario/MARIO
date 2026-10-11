"""Validate the catalogue workflow with fixtures, never real MRIO releases."""

import importlib.util
import json
import sys
from pathlib import Path
from types import SimpleNamespace
import weakref

import pandas as pd
import pytest

from tests._paths import REPO_ROOT

sys.path.insert(0, str(REPO_ROOT / "doc/scripts"))


def _import(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


catalogue = _import("database_catalogue", REPO_ROOT / "doc/scripts/database_catalogue.py")


def _notebook(doc, source, *, prepend=False):
    path = doc / "source/notebooks/parsers/adb/walkthrough_adb.ipynb"
    cells = [{"cell_type": "code", "source": [source]}]
    if prepend:
        cells.insert(0, {"cell_type": "markdown", "source": ["A new introduction"]})
    catalogue.write_json(path, {"cells": cells})
    return path


@pytest.fixture
def setup_job(tmp_path):
    doc = tmp_path / "doc"
    _notebook(doc, 'db = mario.parse_adb(path="/path/to/input.xlsx", year=2020, table="IOT")')
    manifest = catalogue.prepare(doc)
    local = tmp_path / "fixture.xlsx"
    local.touch()
    job = {"id": "fixture-2020", "label": "Fixture 2020", "source": "Test fixture",
           "release": "fixture", "year": 2020, "expected_table": "IOT", "recipe": "adb--01",
           "parameters": {"path": str(local)}}
    return doc, manifest, job


@pytest.mark.parametrize("table,expected", [
    ("IOT", {"Region": 2, "Sector": 3, "Factor of production": 3,
             "Satellite account": 2, "Consumption category": 1}),
    ("SUT", {"Region": 2, "Activity": 2, "Commodity": 2, "Factor of production": 3,
             "Satellite account": 2, "Consumption category": 1}),
])
def test_packaged_sets_and_workbook_roundtrip(table, expected, tmp_path):
    import mario
    from openpyxl import load_workbook

    db = mario.load_test(table)
    recipe = {"guide": "fixture", "notebook": "fixture.ipynb", "cell": 0,
              "id": "fixture--01", "fingerprint": "test", "parser": "load_test"}
    job = {"id": "fixture", "label": "Packaged fixture", "release": "fixture", "year": None}
    profile = catalogue.extract_profile(db, job=job, recipe=recipe, parameters={"path": "/private/input.xlsx"})
    assert profile["counts"] == expected
    assert profile["sets"] == db.get_index("all")
    assert profile["axes"] and profile["units"]
    assert profile["provenance"]["options"]["path"] == "local input supplied"
    profile["sets"]["Region"][0] = "=literal_label"
    path = tmp_path / "sets.xlsx"
    catalogue.export_workbook(profile, path)
    workbook = load_workbook(path)
    assert workbook["Regions"]["B2"].value == "=literal_label"
    assert workbook["Regions"]["B2"].data_type == "s"
    for name, labels in profile["sets"].items():
        sheet = workbook[catalogue.SET_SHEETS[name]]
        assert [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)] == labels
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
    assert workbook["Native axes"].max_row - 1 == sum(len(a["records"]) for a in profile["axes"])
    workbook.close()


@pytest.mark.parametrize("coefficient_input", [False, True])
def test_heterogeneous_region_memberships_and_semantic_extension_axes(coefficient_input):
    productive = pd.MultiIndex.from_tuples(
        [("A", "Sector", "Energy"), ("B", "Sector", "Transport")], names=["Region", "Level", "Item"])
    extension = pd.MultiIndex.from_tuples(
        [("A", "Energy", "CO2")], names=["Region", "Sector", "Satellite account"])
    blocks = {"Z": pd.DataFrame(999, index=productive, columns=productive),
              "E": pd.DataFrame(12345, index=extension, columns=productive)}
    if coefficient_input:
        blocks = {name.lower(): block for name, block in blocks.items()}
    db = SimpleNamespace(
        table_type="IOT", get_index=lambda _: {"Region": pd.Index(["A", "B"]),
                                               "Sector": ["Energy", "Transport"], "Satellite account": ["CO2"]},
        list_matrices=lambda: [*blocks, "X"],
        get_block=lambda name: blocks[name],
        units={"Satellite account": pd.DataFrame({"unit": ["kg"]}, index=extension)},
    )
    recipe = {"guide": "fixture", "notebook": "fixture.ipynb", "cell": 0, "id": "test",
              "fingerprint": "test", "parser": "test"}
    profile = catalogue.extract_profile(db, job={"id": "test", "label": "Test"}, recipe=recipe, parameters={})
    assert {(i["region"], i["label"]) for i in profile["region_items"]} == {("A", "Energy"), ("B", "Transport")}
    assert profile["region_counts"] == [{"region": "A", "set": "Sector", "count": 1},
                                        {"region": "B", "set": "Sector", "count": 1}]
    assert profile["units"][0]["index"] == ["A", "Energy", "CO2"]
    assert profile["units"][0]["index_names"] == ["Region", "Sector", "Satellite account"]
    assert "12345" not in json.dumps(profile)
    assert profile["geography"]["country_memberships"] == []
    assert profile["classifications"]["mappings"] == []


def test_notebook_recipes_are_read_without_execution(tmp_path):
    marker = tmp_path / "must-not-exist"
    path = _notebook(tmp_path, f'open({str(marker)!r}, "w").write("executed")\n'
                     'db = mario.parse_adb(path="/path/to/input")\n'
                     'bridge = mario.parse_from_pymrio(io=io)')
    calls = catalogue.notebook_calls(path)
    assert not marker.exists()
    assert len(calls) == 2
    assert calls[0]["runnable"]
    assert not calls[1]["runnable"]


def test_resolve_tracks_moved_cells_and_rejects_changed_calls(setup_job):
    doc, manifest, job = setup_job
    path = _notebook(doc, 'db = mario.parse_adb(path="/path/to/input.xlsx", year=2020, table="IOT")', prepend=True)
    recipe, options = catalogue.resolve_job(job, manifest, doc=doc)
    assert recipe["cell"] == 1
    assert options["calc_all"] is False and options["download"] is False
    path.write_text(path.read_text().replace("2020", "2021"))
    with pytest.raises(ValueError, match="notebook example changed"):
        catalogue.resolve_job(job, manifest, doc=doc)


def test_missing_input_never_invokes_parser_and_other_jobs_continue(setup_job, tmp_path):
    import mario

    doc, manifest, job = setup_job
    calls = []

    def parse_adb(*, path, year, table, calc_all, download):
        calls.append((calc_all, download))
        return mario.load_test(table)

    parser = SimpleNamespace(parse_adb=parse_adb, __version__="fixture")
    missing = {**job, "id": "missing", "parameters": {"path": str(tmp_path / "missing.xlsx")}}
    output = tmp_path / "private"
    report = catalogue.run_jobs({"jobs": [missing, job]}, manifest, output, parser_module=parser, doc=doc)
    assert [r["status"] for r in report["results"]] == ["failed", "parsed"]
    assert calls == [(False, False)]
    assert not (output / "missing/profile.json").exists()
    assert (output / job["id"] / "sets.xlsx").is_file()
    published = catalogue.publish(output, doc=doc)
    assert len(published["profiles"]) == 1
    assert str(tmp_path) not in json.dumps(published)
    assert "traceback" not in json.dumps(published)
    assert (doc / "source" / published["profiles"][0]["workbook"]).is_file()

    # A failed refresh cannot reuse last run's successful private profile.
    catalogue.write_json(output / "run-report.json", {"results": [{"id": job["id"], "status": "failed", "error": "private"}]})
    assert catalogue.publish(output, doc=doc)["profiles"] == []
    assert not (doc / "source" / published["profiles"][0]["workbook"]).exists()


def test_configuration_conflicts_fail_before_parser(setup_job, tmp_path):
    doc, manifest, job = setup_job
    with pytest.raises(ValueError, match="year conflicts"):
        catalogue.resolve_job({**job, "year": 2021}, manifest, doc=doc)
    with pytest.raises(ValueError, match="expected_table conflicts"):
        catalogue.resolve_job({**job, "expected_table": "SUT"}, manifest, doc=doc)
    with pytest.raises(ValueError, match="unique"):
        catalogue.run_jobs({"jobs": [job, job]}, manifest, tmp_path, parser_module=SimpleNamespace(), doc=doc)
    with pytest.raises(ValueError, match="Invalid profile id"):
        catalogue.run_jobs({"jobs": [{**job, "id": "../escape"}]}, manifest, tmp_path,
                           parser_module=SimpleNamespace(), doc=doc)


def test_failed_job_releases_database_before_reading_the_next(setup_job, tmp_path):
    import mario

    doc, manifest, job = setup_job
    previous = []

    def parse_adb(*, path, year, table, calc_all, download):
        if previous:
            assert previous[0]() is None, "Previous failed database is still in memory"
            return mario.load_test("IOT")
        database = mario.load_test("SUT")
        previous.append(weakref.ref(database))
        return database

    report = catalogue.run_jobs({"jobs": [{**job, "id": "wrong-table"}, job]}, manifest, tmp_path / "private",
                                parser_module=SimpleNamespace(parse_adb=parse_adb), doc=doc)
    assert [r["status"] for r in report["results"]] == ["failed", "parsed"]


def test_incomplete_publish_leaves_existing_catalogue_unchanged(setup_job, tmp_path):
    doc, manifest, job = setup_job
    output = tmp_path / "private"
    path = doc / "source/_static/data/database_profiles.json"
    previous = json.loads(path.read_text())
    previous["profiles"] = [{"id": job["id"], "label": "Old release", "guide": "adb"}]
    catalogue.write_json(path, previous)
    catalogue.write_json(output / "run-report.json", {"results": [{"id": job["id"], "status": "parsed"}]})
    catalogue.write_json(output / job["id"] / "profile.json", {"id": job["id"], "status": "parsed"})
    with pytest.raises(ValueError, match="Incomplete"):
        catalogue.publish(output, doc=doc)
    assert json.loads(path.read_text()) == previous


def test_profile_rendering_is_shared_and_links_work_from_deep_guides(tmp_path):
    pytest.importorskip("sphinx")
    from sphinx.util.osutil import relative_uri

    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    source = tmp_path / "source"
    workbook = "_static/data/database_profiles/fixture.xlsx"
    (source / workbook).parent.mkdir(parents=True)
    (source / workbook).touch()
    app = SimpleNamespace(srcdir=str(source), config=SimpleNamespace(root_doc="index"),
                          builder=SimpleNamespace(get_relative_uri=lambda a, b: relative_uri(a + ".html", b + ".html")))
    profile = {"id": "fixture", "guide": "adb", "status": "parsed", "label": "ADB 2020 <fixture>",
               "table": "IOT", "counts": {"Region": 2, "Sector": 6, "Satellite account": 4}, "workbook": workbook}
    card = extension._summary_html(app, "user_guide/parsers/index", "adb", [profile], compact=True)
    guide = extension._summary_html(app, "notebooks/parsers/adb/walkthrough_adb", "adb", [profile])
    assert "<strong>2</strong> regions" in card and "<strong>6</strong> sectors" in card
    assert "satellite" not in card.lower()
    assert "ADB 2020 &lt;fixture&gt;" in card
    assert workbook not in card
    assert f'href="../../../{workbook}"' in guide
    assert 'href="../../../resources/data-explorer/scope.html#geographical-coverage"' in guide
    assert extension._summary_html(app, "user_guide/parsers/index", "adb", [], compact=True) == ""
    empty_guide = extension._summary_html(app, "notebooks/parsers/adb/walkthrough_adb", "adb", [])
    assert empty_guide == ""


def test_catalogue_dependency_is_registered_before_environment_pickle(tmp_path):
    pytest.importorskip("sphinx")
    from docutils import nodes

    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    dependencies = []
    env = SimpleNamespace(docname="user_guide/parsers/index", note_dependency=dependencies.append)
    app = SimpleNamespace(srcdir=str(tmp_path), env=env)
    tree = nodes.document("", "")
    tree += extension.database_profile_summary("", guide="adb")
    extension._note_dependency(app, tree)
    assert dependencies == [str(tmp_path / extension.CATALOGUE), str(tmp_path / extension.LICENCES)]
    dependencies.clear()
    env.docname = "notebooks/parsers/adb/walkthrough_adb"
    extension._note_dependency(app, nodes.document("", ""))
    assert dependencies == [str(tmp_path / extension.CATALOGUE), str(tmp_path / extension.LICENCES)]


def test_empty_satellite_placeholder_is_retained_but_not_counted():
    profile = {"table": "IOT", "source": "Fixture", "guide": "fixture", "sets": {"Region": ["ITA"], "Sector": ["Energy"], "Satellite account": ["-"]},
               "units": [{"set": "Satellite account", "values": {"unit": "None"}}], "region_items": []}
    result = catalogue.enrich_profile(profile)
    assert result["sets"]["Satellite account"] == ["-"]
    assert result["counts"]["Satellite account"] == 0
    assert result["set_flags"]["Satellite account"]["status"] == "not_loaded"
    profile["units"][0]["values"]["unit"] = "kg"
    assert catalogue.enrich_profile(profile)["counts"]["Satellite account"] == 1


def test_coverage_grouping_compares_labels_not_only_counts():
    pytest.importorskip("sphinx")
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    def profile(identifier, sector):
        return catalogue.enrich_profile({"id": identifier, "table": "IOT", "source": "Fixture", "guide": "fixture",
                                        "sets": {"Region": ["ITA"], "Sector": [sector]}, "units": [], "region_items": []})
    first, same, changed = profile("first", "Energy"), profile("same", "Energy"), profile("changed", "Transport")
    assert first["counts"] == changed["counts"]
    groups = extension.coverage_groups([first, same, changed])
    assert [[p["id"] for p in group] for group in groups] == [["first", "same"], ["changed"]]
    different_region = profile("different-region", "Energy")
    different_region["sets"]["Region"] = ["FRA"]
    catalogue.enrich_profile(different_region)
    assert len(extension.coverage_groups([first, different_region])) == 2


def test_geo_resolution_keeps_all_european_iso2_codes():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    countries = json.loads((REPO_ROOT / "doc/source/_static/data/world-country-map.json").read_text())["countries"]
    profile = {"source": "EXIOBASE", "guide": "exiobase-monetary", "sets": {"Region": ["GR", "GB", "WA"]}}
    result = geography.resolve_geography(profile, countries)
    assert [r.get("iso3") for r in result["regions"][:2]] == ["GRC", "GBR"]
    assert result["unresolved_count"] == 0


def test_adb_aliases_and_aggregates_are_not_fuzzy_country_matches():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    countries = [{"iso3": "CHN", "iso2": "CN", "name": "China", "official_name": "People's Republic of China"}]
    profile = {"source": "Asian Development Bank", "guide": "adb", "sets": {"Region": ["PRC", "RoW", "Unknown"]}}
    result = geography.resolve_geography(profile, countries)
    assert result["individual_countries"] == 1
    assert result["aggregate_count"] == 1 and result["unresolved_count"] == 1
    assert result["regions"][0]["iso3"] == "CHN"
    assert result["regions"][1]["members"] == []


def test_eora_non_iso_and_historical_region_codes_stay_distinct():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    countries = json.loads((REPO_ROOT / "doc/source/_static/data/world-country-map.json").read_text())["countries"]
    profile = {"source": "EORA", "guide": "eora", "sets": {"Region": ["SDS", "SUD", "ANT", "USR"]}}
    result = geography.resolve_geography(profile, countries)
    assert result["individual_countries"] == 2
    assert result["aggregate_count"] == 2 and result["unresolved_count"] == 0
    assert [r.get("iso3") for r in result["regions"][:2]] == ["SSD", "SDN"]
    assert all(r["historical"] and not r["members"] for r in result["regions"][2:])


def test_emerging_historical_groups_are_not_current_countries():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    countries = json.loads((REPO_ROOT / "doc/source/_static/data/world-country-map.json").read_text())["countries"]
    profile = {"source": "EMERGING", "guide": "emerging", "release": "2.2", "year": 2023,
               "sets": {"Region": ["SER", "MNE", "ANT", "TMP"]}}
    result = geography.resolve_geography(profile, countries)
    assert result["individual_countries"] == 2 and result["aggregate_count"] == 2
    assert result["unresolved_count"] == 0
    assert result["regions"][0]["historical"] and not result["regions"][0]["members"]
    assert result["regions"][2]["historical"] and not result["regions"][2]["members"]
    assert result["regions"][3]["iso3"] == "TLS"
    assert not any(r.get("iso3") == "SRB" for r in result["regions"])
    profile["release"] = "unreviewed"
    assert geography.resolve_geography(profile, countries)["unresolved_count"] == 3


def test_country_border_mesh_draws_shared_edges_once():
    generator = _import("country_map_generator", REPO_ROOT / "doc/scripts/generate_country_map.py")
    import re
    shapes = [{"path": "M0,0L1,0L1,1L0,1Z"}, {"path": "M1,0L2,0L2,1L1,1Z"}]
    mesh = generator.border_mesh(shapes)
    edges = []
    for part in re.findall(r"M([^M]+)", mesh):
        closed = part.endswith("Z")
        points = [tuple(map(float, point.split(","))) for point in part.rstrip("Z").split("L")]
        pairs = list(zip(points, points[1:]))
        if closed:
            pairs.append((points[-1], points[0]))
        edges.extend(tuple(sorted(pair)) for pair in pairs)
    assert len(edges) == 7 and len(set(edges)) == 7
    assert edges.count(((1.0, 0.0), (1.0, 1.0))) == 1


def test_parser_cards_group_years_but_keep_economic_variants(tmp_path):
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    profiles = []
    for identifier, year, system, count in [("ixi-2020", 2020, "ixi", 163),
                                           ("ixi-2021", 2021, "ixi", 163),
                                           ("pxp-2020", 2020, "pxp", 200)]:
        profiles.append({"id": identifier, "source": "EXIOBASE", "guide": "exiobase-monetary",
                         "release": "3.11.2", "year": year, "status": "parsed", "table": "IOT",
                         "label": f"EXIOBASE {year} {system}", "counts": {"Region": 49, "Sector": count,
                                                                                   "Satellite account": 808},
                         "coverage_signatures": {"geography": "same-regions", "economic": system},
                         "provenance": {"options": {"system": system}},
                         "workbook": f"_static/{identifier}.xlsx"})
        path=tmp_path/f"_static/{identifier}.xlsx";path.parent.mkdir(exist_ok=True);path.touch()
    app = SimpleNamespace(srcdir=str(tmp_path), config=SimpleNamespace(root_doc="index"),
                          builder=SimpleNamespace(get_relative_uri=lambda a,b: "../" + b + ".html"))
    card=extension._summary_html(app,"user_guide/parsers/index","exiobase-monetary",profiles,compact=True)
    assert card.count("3.11.2") == 1
    assert card.count("<strong>49</strong> regions") == 1
    assert card.count("<strong>163</strong>") == 1 and card.count("<strong>200</strong>") == 1
    assert "ixi" in card and "pxp" in card
    assert "2020" not in card and "2021" not in card and "808" not in card
    material=extension._catalogue_html(app,"resources/material",profiles)
    assert material.count("<tr>") == 2  # One header and one geography/version row.
    assert material.count('class="database-profile-download"') == 3
    assert all(p["workbook"] in material for p in profiles)


def test_guide_summary_follows_existing_intro_and_lists():
    from docutils import nodes
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    top=nodes.section()
    title=nodes.title("", "Parser walkthrough")
    intro=nodes.paragraph("", "This notebook introduces the database.")
    focus=nodes.bullet_list()
    section=nodes.section();section+=nodes.title("", "What this notebook covers")
    top.extend([title,intro,focus,section])
    panel=nodes.container(classes=["database-guide-summary"])
    extension._insert_guide_summary(top,panel)
    assert top.children == [title,intro,focus,panel,section]


def test_exiobase_source_concordance_covers_territories_without_conflicting_with_individuals():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    profile = json.loads((REPO_ROOT / "doc/source/_static/data/database_profiles/exiobase-3112-2024-ixi.json").read_text())
    geography.add_geography(profile)
    geo = profile["geography"]
    membership = {r["iso3"]: r["region"] for r in geo["country_memberships"]}
    assert membership["SOM"] == "WF" and membership["PSE"] == "WM"
    assert membership["GUF"] == "WL" and membership["GUM"] == "WA"
    assert "ITA" not in membership and "TWN" not in membership
    assert len(membership) == len(geo["country_memberships"])
    assert geo["scope"] == "global"
    assert geo["aggregate_mapping_source"]["sheet"] == "Classification (extended)"
    hybrid = json.loads((REPO_ROOT / "doc/source/_static/data/database_profiles/exiobase-hybrid-3318-2011-sut.json").read_text())
    geography.add_geography(hybrid)
    assert hybrid["geography"]["scope"] == "global"
    assert not hybrid["geography"]["country_memberships"]  # Different native set: no guessed group membership.


def test_map_geometry_matches_database_reporting_units_and_preserves_palestine():
    world = json.loads((REPO_ROOT / "doc/source/_static/data/world-country-map.json").read_text())
    assert len([f for f in world["features"] if f["iso3"] == "SOM"]) == 2
    assert not any(f["iso3"] == "SOL" for f in world["features"])
    assert any(f["iso3"] == "PSE" for f in world["features"])
    assert any(c["iso3"] == "PSE" for c in world["countries"])


def test_shared_set_downloads_preserve_verified_years_and_gaps(tmp_path):
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    path = tmp_path / "_static/shared.xlsx"
    path.parent.mkdir();path.touch()
    profile = {"id":"shared", "year":2024, "years":[1995,1996,1998,2024], "release":"3.11.2", "table":"IOT", "workbook":"_static/shared.xlsx", "label":"Shared native sets", "provenance":{"options":{"system":"ixi"}}}
    app = SimpleNamespace(srcdir=str(tmp_path), config=SimpleNamespace(root_doc="index"), builder=SimpleNamespace(get_relative_uri=lambda a,b:"../"+b+".html"))
    links = extension._download_html(app, "resources/material", [profile])
    assert "1995–1996, 1998, 2024" in links
    assert "1995–2024" not in links
    assert "3.11.2 · ixi" in links


def test_parser_release_titles_and_downloads_use_numeric_order(tmp_path):
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    profiles = [{"id":release, "guide":"exiobase-monetary", "release":release,
        "year":2024, "table":"IOT", "source":"EXIOBASE", "coverage_signatures":{"geography":"same"},
        "label":release, "provenance":{"options":{"system":"ixi"}}}
        for release in ["3.10.1","3.8.2","3.11.2","3.9.6","3.9.4"]]
    title = extension._version_title(profiles)
    assert title == "Monetary · 3.11.2 / 3.10.1 / 3.9.6 / 3.9.4 / 3.8.2"
    assert [p["release"] for p in extension.geographic_groups(profiles)[0]] == ["3.11.2","3.10.1","3.9.6","3.9.4","3.8.2"]
    economies = [{**profiles[0], "guide":"adb", "release":f"{n} economies"} for n in [62,74,71,72]]
    assert extension._version_title(economies) == "74 economies / 72 economies / 71 economies / 62 economies"


def test_adb_lac_geography_uses_source_legend_identifiers():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    world = json.loads((REPO_ROOT / "doc/source/_static/data/world-country-map.json").read_text())
    result = geography.resolve_geography({"source":"Asian Development Bank","guide":"adb","sets":{"Region":["PAR","RoLAC","RoW"]}},world["countries"])
    assert result["regions"][0]["iso3"] == "PRY"
    assert result["aggregate_count"] == 2 and result["unresolved_count"] == 0


def test_emerging_shared_nomenclature_requires_an_exact_verified_native_set():
    geography = _import("database_geography", REPO_ROOT / "doc/scripts/database_geography.py")
    countries = json.loads((REPO_ROOT / "doc/source/_static/data/world-country-map.json").read_text())["countries"]
    known = json.loads((REPO_ROOT / "doc/source/_static/data/database_profiles/emerging-22-2023-co2.json").read_text())
    for release in ["1.0", "2.0", "2.2", "E"]:
        profile = {"source": "EMERGING", "guide": "emerging", "release": release, "year": 2018,
                   "sets": {"Region": known["sets"]["Region"]}}
        result = geography.resolve_geography(profile, countries)
        assert result["unresolved_count"] == 0
        assert result["individual_countries"] == 243 and result["aggregate_count"] == 2
        assert not any(r.get("iso3") == "SRB" for r in result["regions"])
        assert result["region_mapping_sources"][0]["native_geography_sha256"] == known["coverage_signatures"]["geography"]
    profile["sets"] = {"Region": ["ANT", "SER", "TMP", "MNE"]}
    assert geography.resolve_geography(profile, countries)["unresolved_count"] == 3


def test_shared_economic_counts_name_every_equivalent_release():
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    profiles = [{"id":release,"source":"EMERGING","guide":"emerging","release":release,"year":2018,
                 "counts":{"Region":245,"Sector":sectors},
                 "coverage_signatures":{"geography":"same","economic":classification}}
                for release,sectors,classification in [("2.2",133,"v2"),("2.0",133,"v2"),("1.0",134,"v1"),("E",146,"e")]]
    metrics=extension._economic_metrics(profiles)
    assert (133,"sectors","2.2 / 2.0") in metrics
    assert (134,"sectors","1.0") in metrics
    assert (146,"sectors","E") in metrics
    assert len(metrics)==3


def test_licence_metadata_covers_guides_and_distinguishes_exiobase_releases():
    extension = _import("mario_database_profiles", REPO_ROOT / "doc/source/_ext/mario_database_profiles.py")
    catalogue=json.loads((REPO_ROOT / "doc/source/_static/data/database_profiles.json").read_text())
    licences=json.loads((REPO_ROOT / "doc/source/_static/data/database_licences.json").read_text())
    assert {g["id"] for g in catalogue["guides"]} == set(licences["guides"])
    rendered=extension._licence_html("exiobase-monetary",licences)
    assert "CC BY-SA 4.0" in rendered and "3.8.2" in rendered
    assert "Custom non-commercial licence" in rendered and "LICENSE.txt" in rendered
    assert "Source licence" in rendered and "Reuse &amp; attribution" in rendered
    assert "Input database licence" in extension._licence_html("custom-database-from-excel",licences)
